import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

file_name = 'Team_H_Detailed_Overdue_Tasks.xlsx'
print(f"Applying UI styles to {file_name}...")

wb = openpyxl.load_workbook(file_name)

# Define standard styles
header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
title_font = Font(size=14, bold=True)
subtitle_font = Font(size=11, italic=True)

# Process Overview Sheet
if 'Overview' in wb.sheetnames:
    ws = wb['Overview']
    # Insert 3 rows at the top
    ws.insert_rows(1, 3)
    
    ws['A1'] = 'Team H - Overdue Tasks by Project'
    ws['A1'].font = title_font
    
    total_projects = len(wb.sheetnames) - 1
    ws['A2'] = f"Generated Live  |  {total_projects} projects with overdue tasks"
    ws['A2'].font = subtitle_font
    
    ws['A3'] = "Each project with overdue tasks has its own sheet (tab) below listing the individual overdue tasks."
    ws['A3'].font = subtitle_font
    
    # Style the headers in Row 4
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=4, column=col)
        cell.fill = header_fill
        cell.font = header_font
        
    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        col_letter = col[3].column_letter # using row 4 for column letter
        for cell in col:
            if cell.row < 4: continue
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[col_letter].width = adjusted_width
        
    ws.freeze_panes = 'A5'

# Process Project Sheets
for sheet_name in wb.sheetnames:
    if sheet_name == 'Overview':
        continue
        
    ws = wb[sheet_name]
    
    # Insert 2 rows at the top
    ws.insert_rows(1, 2)
    
    ws['A1'] = sheet_name # Or use full project name if available, but sheet name works well
    ws['A1'].font = title_font
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ws.max_column)
    ws['A1'].alignment = Alignment(horizontal='left')
    
    # Style headers in Row 3
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=3, column=col)
        cell.fill = header_fill
        cell.font = header_font
        
    # Column widths specifically for tasks
    # ID, Title, Deadline, Person, Status
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 60 # Task Title usually long
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 15
    
    ws.freeze_panes = 'A4'

wb.save(file_name)
print(f"Successfully applied exact original UI styling to {file_name}.")
