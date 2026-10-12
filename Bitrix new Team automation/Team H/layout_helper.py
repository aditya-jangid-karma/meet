import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def build_grouped_sheet(ws, tasks_data):
    # Group tasks by Client
    # tasks_data is a list of dicts: Type, Client Code, Client, Task ID, Task Title, Deadline, Responsible Person, Status
    
    # Sort first by Client Name, then by Type (Overdue, Due Today, Due Tomorrow), then Deadline
    df = pd.DataFrame(tasks_data)
    if df.empty:
        return
        
    # We want a specific sort order for Type
    type_order = {'overdue': 0, 'today': 1, 'tomorrow': 2}
    df['Type_Sort'] = df['Type'].map(type_order)
    df = df.sort_values(by=['Client', 'Type_Sort', 'Deadline'])
    
    headers = ['Type', 'Task ID', 'Task Title', 'Deadline', 'Responsible Person', 'Status']
    header_fill = PatternFill(start_color="1A3B69", end_color="1A3B69", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    group_font = Font(size=14, bold=True)
    alt_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                         top=Side(style='thin'), bottom=Side(style='thin'))
                         
    row_idx = 1
    
    for client, group in df.groupby('Client', sort=False):
        # 1. Big header row
        code = group.iloc[0]['Client Code']
        ws.cell(row=row_idx, column=1, value=f"{code} {client}")
        ws.cell(row=row_idx, column=1).font = group_font
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=len(headers))
        row_idx += 1
        
        # 2. Column headers
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=h)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal='center')
        row_idx += 1
        
        # 3. Tasks
        for _, row in group.iterrows():
            vals = [
                row['Type'].capitalize(),
                row['Task ID'],
                row['Task Title'],
                row['Deadline'],
                row['Responsible Person'],
                row['Status']
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.border = thin_border
            row_idx += 1
            
        # Blank row
        row_idx += 1
        
    # Column widths
    ws.column_dimensions['A'].width = 15 # Type
    ws.column_dimensions['B'].width = 12 # Task ID
    ws.column_dimensions['C'].width = 45 # Task Title
    ws.column_dimensions['D'].width = 15 # Deadline
    ws.column_dimensions['E'].width = 25 # Responsible Person
    ws.column_dimensions['F'].width = 15 # Status

