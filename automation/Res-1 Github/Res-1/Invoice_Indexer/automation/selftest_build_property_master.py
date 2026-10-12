"""Offline tests for build_property_master.py (real-shape sheets). Builds tiny
in-memory workbooks mimicking the REAL Yardi (coded table) and Intacct (banner
+ spacer report) sheets. Never opens the real Vendor&PropertyMaster.xlsx. Run:
  python automation\\selftest_build_property_master.py
"""
import csv
import os
import tempfile
import unittest

from openpyxl import Workbook

from build_property_master import (
    find_sheet, rows_from_yardi_sheet, rows_from_intacct_sheet,
    build_rows, write_master, FIELDNAMES,
)


class FindSheet(unittest.TestCase):
    def test_exact_name_matches(self):
        wb = Workbook()
        wb.create_sheet(title="Foo")
        self.assertEqual(find_sheet(wb, "Foo").title, "Foo")

    def test_trailing_space_and_case_are_forgiven(self):
        wb = Workbook()
        wb.create_sheet(title="HOC - YARDI ")   # real sheet has a trailing space
        self.assertEqual(find_sheet(wb, "HOC - YARDI").title, "HOC - YARDI ")
        self.assertEqual(find_sheet(wb, "hoc - yardi").title, "HOC - YARDI ")

    def test_missing_sheet_raises_with_available_names_listed(self):
        wb = Workbook()
        wb.create_sheet(title="Foo")
        with self.assertRaises(ValueError) as ctx:
            find_sheet(wb, "Bar")
        self.assertIn("Foo", str(ctx.exception))


def _yardi_sheet(wb):
    """Header row 1 -- 'Property' (code) and 'Property Name', plus other
    columns, several units per property, same code every time -- exactly the
    real HOC - YARDI shape."""
    ws = wb.create_sheet(title="Yardi")
    ws.append(["Property", "Property Name", "Unit", "City"])
    ws.append(["414-460", "Fairfax Court Apts", 101, "Chevy Chase"])
    ws.append(["414-460", "Fairfax Court Apts", 102, "Chevy Chase"])
    ws.append(["439-746", "Barclay One", 201, "Somewhere"])
    return ws


def _intacct_sheet(wb):
    """Banner rows, header on row 6, blank spacer row between every unit --
    exactly the real 'All Units - Intacct' shape."""
    ws = wb.create_sheet(title="Intacct")
    ws.append(["All Units Report"])
    ws.append([None] * 13 + ["Printed on: 5/1/26"])
    ws.append(["Community = All"])
    ws.append([])
    ws.append([])
    ws.append(["Property Name", "Apt Nbr", None, "Mailing Address"])
    ws.append([])                                    # spacer
    ws.append(["11Hundred", "101", None, "1100 Eastern Ave"])
    ws.append([])                                    # spacer
    ws.append(["11Hundred", "102", None, "1100 Eastern Ave"])
    ws.append([])                                    # spacer
    ws.append(["Oakwood", "201", None, "9 Test Ave"])
    return ws


class RowsFromYardiSheet(unittest.TestCase):
    def test_dedupes_and_captures_code(self):
        wb = Workbook()
        ws = _yardi_sheet(wb)
        rows = rows_from_yardi_sheet(ws)
        by_name = {r["Property Name"]: r for r in rows}
        self.assertEqual(len(rows), 2)   # 3 unit rows -> 2 distinct properties
        self.assertEqual(by_name["Fairfax Court Apts"]["Property Code"], "414-460")
        self.assertEqual(by_name["Fairfax Court Apts"]["Package"], "YARDI")
        self.assertEqual(by_name["Barclay One"]["Property Code"], "439-746")


class RowsFromIntacctSheet(unittest.TestCase):
    def test_finds_header_past_banner_rows_dedupes_and_skips_blanks(self):
        wb = Workbook()
        ws = _intacct_sheet(wb)
        rows = rows_from_intacct_sheet(ws)
        names = sorted(r["Property Name"] for r in rows)
        self.assertEqual(names, ["11Hundred", "Oakwood"])
        for r in rows:
            self.assertEqual(r["Package"], "INTACCT")
            self.assertEqual(r["Property Code"], "")   # Intacct has no codes


class BuildRows(unittest.TestCase):
    def test_combines_both_with_correct_packages_and_codes(self):
        wb = Workbook()
        _yardi_sheet(wb)
        _intacct_sheet(wb)
        rows = build_rows(wb, "Yardi", "Intacct")
        packages = {r["Property Name"]: r["Package"] for r in rows}
        self.assertEqual(packages["Fairfax Court Apts"], "YARDI")
        self.assertEqual(packages["11Hundred"], "INTACCT")


class WriteMaster(unittest.TestCase):
    def test_writes_property_code_column_blank_for_intacct(self):
        rows = [{"Property Name": "Fairfax Court Apts", "Package": "YARDI",
                 "Property Code": "414-460"},
                {"Property Name": "11Hundred", "Package": "INTACCT",
                 "Property Code": ""}]
        fd, path = tempfile.mkstemp(suffix=".csv")
        os.close(fd)
        try:
            write_master(rows, path)
            with open(path, newline="", encoding="utf-8-sig") as f:
                got = list(csv.DictReader(f))
            self.assertEqual(got[0]["Property Code"], "414-460")
            self.assertEqual(got[1]["Property Code"], "")
            self.assertEqual(list(got[0].keys()), FIELDNAMES)
        finally:
            os.remove(path)


if __name__ == "__main__":
    unittest.main()
