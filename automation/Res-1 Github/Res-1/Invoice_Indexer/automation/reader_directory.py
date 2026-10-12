"""Match a vendor name from Qwen against the vendors already recorded in
invoices.xlsx. Until a real vendor master exists, the sheet IS the master
(see the spec, D10). This file reads the sheet and the model's ANSWER only --
never a PDF.
"""
from openpyxl import load_workbook


def _normalise(name):
    """Lower-case, trim, and collapse inner spaces so 'A  B ' == 'a b'."""
    return " ".join(str(name).lower().split())


def load_known_vendors(xlsx_path):
    """Return {normalised vendor -> exact recorded spelling} from the sheet's
    Vendor column. An empty or vendor-less sheet yields an empty dict, which is
    correct for the bootstrap phase (everything is a 'new' vendor)."""
    workbook = load_workbook(xlsx_path, data_only=True)
    sheet = (workbook["Invoices"] if "Invoices" in workbook.sheetnames
             else workbook.active)

    vendor_col = None
    for cell in sheet[1]:
        if cell.value and _normalise(cell.value) == "vendor":
            vendor_col = cell.column
            break
    if vendor_col is None:
        return {}

    known = {}
    for row in sheet.iter_rows(min_row=2):
        value = row[vendor_col - 1].value
        if value not in (None, ""):
            known[_normalise(value)] = str(value)
    return known


def match_vendor(name, known):
    """Canonical spelling if this vendor is already recorded, else None."""
    return known.get(_normalise(name))
