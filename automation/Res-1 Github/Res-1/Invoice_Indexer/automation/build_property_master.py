"""build_property_master.py -- turn the property data in
Vendor&PropertyMaster.xlsx into masters\\properties.csv.

PLAIN ENGLISH:
    Two sheets hold the real, canonical property lists (a THIRD sheet,
    "Property Names", is a human-typed list that may not match the software's
    exact spelling, so it is NOT used here):

      - The YARDI sheet ("HOC - YARDI ", note the trailing space in the real
        file) is a clean table: one row per UNIT, with the property's NAME in
        one column and its CODE in another. Typing the CODE into Bottomline's
        Property box loads the property automatically -- exactly like a
        vendor's ERP Vendor ID -- so we capture it.
      - The INTACCT sheet ("All Units - Intacct") is a printed REPORT, not a
        clean table: a few banner rows, blank spacer rows between every unit,
        and the real header several rows down. There is no property code here.

    Both sheets repeat the same property once per unit. This script finds the
    real header row on each sheet (scanning down, not assuming row 1), reads
    every unit row, and DEDUPES down to one row per distinct property name,
    tagging YARDI rows with their code and INTACCT rows with none.

    NEVER writes to the workbook. Re-run whenever the property lists change.

Run (sheet names confirmed with the team first -- use the EXACT names from
workbook.sheetnames; matching below forgives only leading/trailing spaces and
case, so "HOC - YARDI" on the command line still finds the real
"HOC - YARDI " sheet):
    .venv\\Scripts\\python.exe automation\\build_property_master.py ^
        --xlsx "Vendor&PropertyMaster.xlsx" ^
        --yardi-sheet "HOC - YARDI" ^
        --intacct-sheet "All Units - Intacct"

Depends only on openpyxl + the Python standard library.
"""
import argparse
import csv
import os
import tempfile
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent
DEFAULT_OUT = PROJECT / "masters" / "properties.csv"

FIELDNAMES = ["Property Name", "Package", "Property Code"]


def find_sheet(workbook, wanted_name):
    """Find a sheet by name, forgiving stray leading/trailing spaces and case
    -- the real workbook has a sheet literally named 'HOC - YARDI ' (trailing
    space), which an exact match would silently miss."""
    if wanted_name in workbook.sheetnames:
        return workbook[wanted_name]
    target = wanted_name.strip().casefold()
    for name in workbook.sheetnames:
        if name.strip().casefold() == target:
            return workbook[name]
    raise ValueError(
        f"no sheet named '{wanted_name}' (ignoring case/spacing). "
        f"Sheets in this workbook: {workbook.sheetnames}")


def _find_header_row(ws, required_headers, max_scan_rows=20):
    """Scan the first `max_scan_rows` rows for one that contains EVERY header
    in `required_headers` (exact text match, case-insensitive, trimmed).
    Returns (row_number, {header -> 0-based column index}). Raises ValueError
    if no such row is found within the scan window -- this is a REPORT sheet
    (Intacct), not a plain table, so the header is not always row 1."""
    wanted = {h.strip().casefold() for h in required_headers}
    for row_number, values in enumerate(
            ws.iter_rows(min_row=1, max_row=max_scan_rows, values_only=True), start=1):
        found = {}
        for col_idx, cell in enumerate(values or ()):
            if cell in (None, ""):
                continue
            text = str(cell).strip()
            if text.casefold() in wanted:
                found[text.casefold()] = col_idx
        if len(found) == len(wanted):
            columns = {h: found[h.strip().casefold()] for h in required_headers}
            return row_number, columns
    raise ValueError(
        f"could not find a row with all of {list(required_headers)} in the "
        f"first {max_scan_rows} rows of sheet '{ws.title}'")


def rows_from_yardi_sheet(ws, name_header="Property Name", code_header="Property"):
    """One row per DISTINCT property name from the Yardi unit-level sheet,
    tagged Package=YARDI, with the property's CODE captured from the FIRST
    unit row seen for that name (the code is the same on every unit of one
    property, confirmed against the real sheet)."""
    header_row, cols = _find_header_row(ws, [name_header, code_header])
    name_idx = cols[name_header]
    code_idx = cols[code_header]

    seen = {}   # normalised name -> row dict, preserves first-seen order
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        name = values[name_idx] if name_idx < len(values) else None
        if name in (None, "") or not str(name).strip():
            continue
        name = str(name).strip()
        key = name.casefold()
        if key in seen:
            continue
        code = values[code_idx] if code_idx < len(values) else None
        code = str(code).strip() if code not in (None, "") else ""
        seen[key] = {"Property Name": name, "Package": "YARDI",
                     "Property Code": code}
    return list(seen.values())


def rows_from_intacct_sheet(ws, name_header="Property Name"):
    """One row per DISTINCT property name from the Intacct report sheet,
    tagged Package=INTACCT, no code (Intacct has none). The sheet is a
    printed report with banner and blank spacer rows -- _find_header_row
    locates the real header, and blank names are simply skipped."""
    header_row, cols = _find_header_row(ws, [name_header])
    name_idx = cols[name_header]

    seen = {}
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        name = values[name_idx] if name_idx < len(values) else None
        if name in (None, "") or not str(name).strip():
            continue
        name = str(name).strip()
        key = name.casefold()
        if key in seen:
            continue
        seen[key] = {"Property Name": name, "Package": "INTACCT",
                     "Property Code": ""}
    return list(seen.values())


def build_rows(workbook, yardi_sheet, intacct_sheet,
               name_header="Property Name", yardi_code_header="Property"):
    """Combine both sheets: Yardi rows (with codes) + Intacct rows (no codes)."""
    yardi_ws = find_sheet(workbook, yardi_sheet)
    intacct_ws = find_sheet(workbook, intacct_sheet)
    return (rows_from_yardi_sheet(yardi_ws, name_header, yardi_code_header)
            + rows_from_intacct_sheet(intacct_ws, name_header))


def write_master(rows, csv_path):
    """Write properties.csv atomically (temp file + os.replace). Always
    Property Name, Package, Property Code (blank for Intacct rows)."""
    directory = os.path.dirname(os.path.abspath(csv_path)) or "."
    os.makedirs(directory, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".properties_tmp_", dir=directory)
    try:
        with os.fdopen(fd, "w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=FIELDNAMES)
            writer.writeheader()
            for row in rows:
                writer.writerow({k: row.get(k, "") for k in FIELDNAMES})
        os.replace(tmp, csv_path)
    except Exception:
        if os.path.exists(tmp):
            os.remove(tmp)
        raise


def main():
    parser = argparse.ArgumentParser(
        description="Build masters\\properties.csv from the Yardi (coded) "
                     "and Intacct (report) property sheets in "
                     "Vendor&PropertyMaster.xlsx.")
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--yardi-sheet", required=True)
    parser.add_argument("--intacct-sheet", required=True)
    parser.add_argument("--name-header", default="Property Name")
    parser.add_argument("--yardi-code-header", default="Property")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    if not args.xlsx.exists():
        raise SystemExit(f"STOPPED: no workbook at {args.xlsx}")

    workbook = load_workbook(args.xlsx, read_only=True, data_only=True)
    try:
        rows = build_rows(workbook, args.yardi_sheet, args.intacct_sheet,
                          args.name_header, args.yardi_code_header)
    finally:
        workbook.close()

    write_master(rows, args.out)
    yardi_n = sum(1 for r in rows if r["Package"] == "YARDI")
    intacct_n = sum(1 for r in rows if r["Package"] == "INTACCT")
    print(f"Wrote {len(rows)} property row(s) to {args.out} "
          f"({yardi_n} YARDI, {intacct_n} INTACCT)")


if __name__ == "__main__":
    main()
