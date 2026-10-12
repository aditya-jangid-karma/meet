"""Offline tests for the Qwen auto-fill reader.

No network, no Poppler, no Bottomline, no live Qwen. Pure logic only.
Run:  python automation\\selftest_reader.py
"""
import unittest

import io
from PIL import Image

from reader_directory import _normalise, match_vendor
from reader_qwen import build_schema, build_prompt, extract_invoice
from reader_pdf import render_pdf
from vendor_lookup import Lookup
from invoice_reader import (append_invoice, _key, existing_image_names,
                            last_data_row, _to_mmddyyyy)
from master_index import build_index, extra_for_confident_match


class VendorMatch(unittest.TestCase):
    def setUp(self):
        # A tiny stand-in for load_known_vendors()'s output.
        self.known = {_normalise("Ferguson Facility Supply"):
                      "Ferguson Facility Supply"}

    def test_exact_known_vendor_returns_canonical(self):
        self.assertEqual(
            match_vendor("ferguson facility supply", self.known),
            "Ferguson Facility Supply")

    def test_case_and_space_differences_still_match(self):
        self.assertEqual(
            match_vendor("  FERGUSON   Facility Supply ", self.known),
            "Ferguson Facility Supply")

    def test_unknown_vendor_returns_none(self):
        self.assertIsNone(match_vendor("Brand New Vendor LLC", self.known))


class QwenExtract(unittest.TestCase):
    GOOD = ('{"vendor":"Acme LLC","invoice_number":"007458",'
            '"invoice_date":"2026-08-01","due_date":null,"amount":"1234.56"}')

    def _post_returning(self, *responses):
        """Fake _post: yields each canned Ollama envelope in turn."""
        calls = {"n": 0}
        seq = list(responses)

        def fake_post(host, payload):
            i = min(calls["n"], len(seq) - 1)
            calls["n"] += 1
            return {"response": seq[i]}
        return fake_post

    def test_schema_marks_due_date_optional(self):
        schema = build_schema()
        self.assertIn("due_date", schema["properties"])
        self.assertNotIn("due_date", schema["required"])

    def test_schema_marks_property_optional(self):
        schema = build_schema()
        self.assertIn("property", schema["properties"])
        self.assertNotIn("property", schema["required"])

    def test_good_json_first_try(self):
        out = extract_invoice("http://x", ["b64"],
                              post=self._post_returning(self.GOOD))
        self.assertEqual(out["invoice_number"], "007458")
        self.assertIsNone(out["due_date"])

    def test_bad_then_good_uses_retry(self):
        out = extract_invoice("http://x", ["b64"],
                              post=self._post_returning("not json", self.GOOD))
        self.assertEqual(out["vendor"], "Acme LLC")

    def test_two_failures_returns_none(self):
        out = extract_invoice("http://x", ["b64"],
                              post=self._post_returning("nope", "still nope"))
        self.assertIsNone(out)

    def test_extract_invoice_still_works_with_alias_examples(self):
        # The nudge is just extra prompt text -- extract_invoice's own
        # behaviour (parse, retry, return dict) must not change.
        out = extract_invoice(
            "http://x", ["b64"], post=self._post_returning(self.GOOD),
            alias_examples=[("Fergusen", "Ferguson Facility Supply")])
        self.assertEqual(out["vendor"], "Acme LLC")


class PromptNudge(unittest.TestCase):
    """build_prompt() with/without alias_examples -- the 'nudge' feature."""

    def test_none_matches_no_args(self):
        self.assertEqual(build_prompt(None), build_prompt())

    def test_empty_list_matches_no_args(self):
        self.assertEqual(build_prompt([]), build_prompt())

    def test_no_examples_has_no_corrections_block(self):
        self.assertNotIn("KNOWN CORRECTIONS", build_prompt())

    def test_examples_appended_with_raw_and_correct_and_delimiter(self):
        prompt = build_prompt([("Fergusen", "Ferguson Facility Supply")])
        self.assertIn("KNOWN CORRECTIONS", prompt)
        self.assertIn("Fergusen", prompt)
        self.assertIn("Ferguson Facility Supply", prompt)
        # The original instructions must still be present, unchanged.
        self.assertTrue(prompt.startswith(build_prompt()))

    def test_multiple_examples_each_get_their_own_line(self):
        prompt = build_prompt([
            ("Fergusen", "Ferguson Facility Supply"),
            ("Ace Plumbng", "Ace Plumbing Inc"),
        ])
        self.assertIn("Fergusen", prompt)
        self.assertIn("Ace Plumbng", prompt)
        self.assertIn("Ace Plumbing Inc", prompt)


def _fake_convert_factory(page_count):
    """Return a stand-in for pdf2image.convert_from_path yielding blank pages."""
    def fake_convert(path, dpi=200, first_page=None, last_page=None,
                     poppler_path=None):
        last = min(last_page or page_count, page_count)
        return [Image.new("RGB", (5, 5), "white")
                for _ in range(first_page or 1, last + 1)]
    return fake_convert


class PdfRender(unittest.TestCase):
    def test_returns_png_bytes_per_page(self):
        out = render_pdf("x.pdf", _convert=_fake_convert_factory(2))
        self.assertEqual(len(out), 2)
        # Each item is real PNG bytes.
        Image.open(io.BytesIO(out[0])).verify()

    def test_caps_at_max_pages(self):
        out = render_pdf("x.pdf", max_pages=3, _convert=_fake_convert_factory(9))
        self.assertEqual(len(out), 3)


HEADERS = ["Image Name", "Vendor", "Property", "Invoice Number", "Invoice Date",
           "Invoice Due Date", "Amount", "Integration Package", "Status",
           "Qwen Vendor (raw)", "Qwen Status",
           "Qwen Property (raw)", "Property Status"]


def _columns(sheet):
    return {_key(c.value): c.column for c in sheet[1] if c.value}


class RowWriting(unittest.TestCase):
    def _fresh_sheet(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Invoices"
        ws.append(HEADERS)
        return wb, ws

    def _append_at_end(self, ws, fields, image_name):
        cols = _columns(ws)
        row = last_data_row(ws, cols) + 1
        append_invoice(ws, cols, fields, image_name, row)
        return row

    def test_append_writes_fields_and_leaves_blanks(self):
        wb, ws = self._fresh_sheet()
        fields = {"vendor": "Acme LLC", "invoice_number": "007458",
                  "invoice_date": "2026-08-01", "due_date": None,
                  "amount": "1234.56"}
        r = self._append_at_end(ws, fields, image_name="skm_0001")

        got = {_key(ws.cell(row=1, column=c.column).value): ws.cell(
            row=r, column=c.column).value for c in ws[1] if c.value}
        self.assertEqual(got["image_name"], "skm_0001")
        self.assertEqual(got["invoice_number"], "007458")   # text, zero kept
        self.assertEqual(got["invoice_date"], "08/01/2026")  # forced MM/DD/YYYY
        self.assertEqual(got["invoice_due_date"], None)      # blank, not printed
        self.assertIn(got["property"], (None, ""))           # none read -> blank
        self.assertIn(got["integration_package"], (None, ""))

    def test_append_writes_property_when_qwen_reads_one(self):
        wb, ws = self._fresh_sheet()
        fields = {"vendor": "Acme LLC", "invoice_number": "007458",
                  "invoice_date": "2026-08-01", "due_date": None,
                  "amount": "1234.56", "property": "4200 Heather Hill Rd"}
        r = self._append_at_end(ws, fields, image_name="skm_0002")

        got = {_key(ws.cell(row=1, column=c.column).value): ws.cell(
            row=r, column=c.column).value for c in ws[1] if c.value}
        self.assertEqual(got["property"], "4200 Heather Hill Rd")  # raw, written

    def test_append_writes_qwen_helper_columns_when_present(self):
        wb, ws = self._fresh_sheet()
        fields = {"vendor": "Ferguson Facility Supply",
                  "invoice_number": "007458", "invoice_date": "2026-08-01",
                  "due_date": None, "amount": "1234.56"}
        cols = _columns(ws)
        row = last_data_row(ws, cols) + 1
        append_invoice(ws, cols, fields, "skm_0003", row,
                       raw_vendor="ferguson facilty supply",
                       qwen_status="NEEDS_REVIEW")

        got = {_key(ws.cell(row=1, column=c.column).value): ws.cell(
            row=row, column=c.column).value for c in ws[1] if c.value}
        self.assertEqual(got["qwen_vendor_raw"], "ferguson facilty supply")
        self.assertEqual(got["qwen_status"], "NEEDS_REVIEW")

    def test_lookup_read_ok_writes_canonical_vendor_and_raw_and_status(self):
        # Drive this via a tiny in-memory master, the same as vendor_lookup's
        # own selftest, then feed the Match straight into append_invoice --
        # this is exactly what main() does for each invoice.
        master_rows = [{"Vendor Name": "Ferguson Facility Supply",
                        "Address": "123 Main St, Lenexa, KS 66214"}]
        lookup = Lookup(master_rows=master_rows)
        raw_vendor = "  FERGUSON   Facility Supply "   # case/space mangled
        m = lookup.match(raw_vendor)
        self.assertEqual(m.status, "READ_OK")

        wb, ws = self._fresh_sheet()
        fields = {"vendor": m.value, "invoice_number": "1",
                  "invoice_date": "2026-08-01", "due_date": None,
                  "amount": "10"}
        cols = _columns(ws)
        row = last_data_row(ws, cols) + 1
        append_invoice(ws, cols, fields, "skm_0004", row,
                       raw_vendor=raw_vendor, qwen_status=m.status)

        got = {_key(ws.cell(row=1, column=c.column).value): ws.cell(
            row=row, column=c.column).value for c in ws[1] if c.value}
        self.assertEqual(got["vendor"], "Ferguson Facility Supply")  # canonical
        self.assertEqual(got["qwen_vendor_raw"], raw_vendor)          # raw text kept
        self.assertEqual(got["qwen_status"], "READ_OK")

    def test_lookup_needs_review_keeps_raw_vendor_and_flags_status(self):
        master_rows = [{"Vendor Name": "Ferguson Facility Supply",
                        "Address": "123 Main St, Lenexa, KS 66214"}]
        lookup = Lookup(master_rows=master_rows)
        raw_vendor = "Completely Different Vendor LLC"   # nothing close
        m = lookup.match(raw_vendor)
        self.assertEqual(m.status, "NEEDS_REVIEW")

        wb, ws = self._fresh_sheet()
        fields = {"vendor": m.value, "invoice_number": "2",
                  "invoice_date": "2026-08-01", "due_date": None,
                  "amount": "20"}
        cols = _columns(ws)
        row = last_data_row(ws, cols) + 1
        append_invoice(ws, cols, fields, "skm_0005", row,
                       raw_vendor=raw_vendor, qwen_status=m.status)

        got = {_key(ws.cell(row=1, column=c.column).value): ws.cell(
            row=row, column=c.column).value for c in ws[1] if c.value}
        self.assertEqual(got["vendor"], raw_vendor)          # RAW kept, not auto-filled
        self.assertEqual(got["qwen_vendor_raw"], raw_vendor)
        self.assertEqual(got["qwen_status"], "NEEDS_REVIEW")

    def test_invoice_number_is_text_format(self):
        wb, ws = self._fresh_sheet()
        r = self._append_at_end(
            ws, {"vendor": "A", "invoice_number": "007458",
                 "invoice_date": "07/28/2026", "due_date": None, "amount": "1"},
            image_name="x")
        col = _columns(ws)["invoice_number"]
        self.assertEqual(ws.cell(row=r, column=col).number_format, "@")

    def test_existing_image_names_normalises(self):
        wb, ws = self._fresh_sheet()
        ws.append(["SKM_0001", "A", "", "1", "07/28/2026", "", "1", "", ""])
        self.assertIn("skm_0001", existing_image_names(ws))

    def test_last_data_row_ignores_phantom_maxrow(self):
        """openpyxl's max_row is inflated by empty-but-formatted cells far down
        the sheet (the bug that stranded the first live rows). last_data_row
        must return the TRUE last data row so the next row lands right after
        the visible data, not hundreds of rows below."""
        wb, ws = self._fresh_sheet()
        ws.append(["skm_0001", "A", "", "1", "07/28/2026", "", "1", "", ""])
        ws.cell(row=200, column=9)          # phantom empty cell inflates max_row
        self.assertEqual(ws.max_row, 200)
        self.assertEqual(last_data_row(ws, _columns(ws)), 2)

    def test_append_lands_right_after_visible_data(self):
        wb, ws = self._fresh_sheet()
        ws.append(["skm_0001", "A", "", "1", "07/28/2026", "", "1", "", ""])
        ws.cell(row=200, column=9)          # phantom inflates max_row to 200
        r = self._append_at_end(
            ws, {"vendor": "B", "invoice_number": "2", "invoice_date": "d",
                 "due_date": None, "amount": "2"}, image_name="skm_0002")
        self.assertEqual(r, 3)              # NOT 201


class DateFormat(unittest.TestCase):
    def test_various_inputs_become_mmddyyyy(self):
        for raw, want in [("2026-07-28", "07/28/2026"),
                          ("07/28/2026", "07/28/2026"),
                          ("July 28, 2026", "07/28/2026"),
                          ("28 July 2026", "07/28/2026"),
                          ("7/8/2026", "07/08/2026")]:
            self.assertEqual(_to_mmddyyyy(raw), want)

    def test_unparseable_is_kept_not_invented(self):
        self.assertEqual(_to_mmddyyyy("see attached"), "see attached")

    def test_blank_stays_blank(self):
        self.assertIsNone(_to_mmddyyyy(None))
        self.assertEqual(_to_mmddyyyy(""), "")


class PropertyWriting(unittest.TestCase):
    """The reader's property side: canonical name + package on a confident
    match; raw + blank package on review. Mirrors the vendor tests above."""

    PROP_ROWS = [{"Property Name": "Heather Hill", "Package": "YARDI"},
                 {"Property Name": "Oakwood", "Package": "INTACCT"}]

    def _fresh_sheet(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Invoices"
        ws.append(HEADERS)
        return wb, ws

    def _got(self, ws, row):
        return {_key(ws.cell(row=1, column=c.column).value):
                ws.cell(row=row, column=c.column).value
                for c in ws[1] if c.value}

    def test_confident_property_writes_canonical_and_package(self):
        lookup = Lookup(master_rows=self.PROP_ROWS, name_col="Property Name")
        pkg_index = build_index(self.PROP_ROWS, "Property Name", "Package")
        raw_property = "  HEATHER   hill "
        m = lookup.match(raw_property)
        self.assertEqual(m.status, "READ_OK")
        package = extra_for_confident_match(m, pkg_index)

        wb, ws = self._fresh_sheet()
        fields = {"vendor": "Acme", "invoice_number": "1",
                  "invoice_date": "2026-08-01", "due_date": None,
                  "amount": "10", "property": m.value}
        cols = _columns(ws)
        row = last_data_row(ws, cols) + 1
        append_invoice(ws, cols, fields, "skm_p1", row,
                       raw_property=raw_property, property_status=m.status,
                       integration_package=package)

        got = self._got(ws, row)
        self.assertEqual(got["property"], "Heather Hill")        # canonical
        self.assertEqual(got["integration_package"], "YARDI")    # package hint
        self.assertEqual(got["qwen_property_raw"], raw_property)
        self.assertEqual(got["property_status"], "READ_OK")

    def test_review_property_keeps_raw_and_leaves_package_blank(self):
        lookup = Lookup(master_rows=self.PROP_ROWS, name_col="Property Name")
        pkg_index = build_index(self.PROP_ROWS, "Property Name", "Package")
        raw_property = "Somewhere Not In The Master"
        m = lookup.match(raw_property)
        self.assertEqual(m.status, "NEEDS_REVIEW")
        package = extra_for_confident_match(m, pkg_index)   # "" -> blank

        wb, ws = self._fresh_sheet()
        fields = {"vendor": "Acme", "invoice_number": "2",
                  "invoice_date": "2026-08-01", "due_date": None,
                  "amount": "20", "property": m.value}
        cols = _columns(ws)
        row = last_data_row(ws, cols) + 1
        append_invoice(ws, cols, fields, "skm_p2", row,
                       raw_property=raw_property, property_status=m.status,
                       integration_package=package)

        got = self._got(ws, row)
        self.assertEqual(got["property"], raw_property)          # raw kept
        self.assertIn(got["integration_package"], (None, ""))    # blank on review
        self.assertEqual(got["property_status"], "NEEDS_REVIEW")


if __name__ == "__main__":
    unittest.main()
