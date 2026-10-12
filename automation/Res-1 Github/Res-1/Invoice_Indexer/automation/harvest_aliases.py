"""Turn a human's spreadsheet corrections into a growing "alias" file.

PLAIN ENGLISH -- what this script does:

  When a reviewer fixes the "Vendor" cell in invoices.xlsx (because Qwen read
  the vendor name wrong, or slightly differently than the master spells it),
  that fix is a lesson we should not have to re-learn by hand every time. This
  script scans the "Invoices" sheet, finds rows where "Vendor" no longer
  matches "Qwen Vendor (raw)" (the reader's original guess, never hand-edited),
  and figures out what kind of correction happened:

    - Corrected to a name that IS already in masters\\vendors.csv
        -> a genuine alias: "next time you see this raw text, it really means
           this vendor". Recorded in masters\\aliases.csv (raw,correct,count,
           last_seen) so vendor_lookup.py's alias step picks it up automatically
           on future reads.
    - Corrected to a name that is NOT in the vendor master
        -> a brand-new vendor. Per design rule D5 ("new vendors are added to
           the master only with a human's OK"), this script never invents a
           new alias or touches the master for this case -- it just reports
           the name so a human (Puneet/Manish) can decide whether to add it.

  Like the reader (invoice_reader.py --dry-run), this script defaults to
  PREVIEW mode: it tells you exactly what it WOULD change and writes nothing.
  You only get a real update to masters\\aliases.csv by adding --apply.

Run (preview -- safe, nothing written):
    .venv\\Scripts\\python.exe automation\\harvest_aliases.py --xlsx PATH\\TO\\invoices.xlsx

Run for real (updates masters\\aliases.csv):
    .venv\\Scripts\\python.exe automation\\harvest_aliases.py --xlsx PATH\\TO\\invoices.xlsx --apply

This file only depends on:
  - openpyxl   (already installed) to read the Invoices sheet
  - csv, os, tempfile, argparse, sys, datetime  -- Python standard library
  - vendor_lookup._normalise / load_master      -- reused, not re-implemented
"""
import argparse
import csv
import os
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from vendor_lookup import _normalise, load_master

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent
MASTERS_DIR = PROJECT / "masters"
VENDORS_CSV = MASTERS_DIR / "vendors.csv"
ALIASES_CSV = MASTERS_DIR / "aliases.csv"

ALIAS_FIELDNAMES = ["raw", "correct", "count", "last_seen"]


def _key(header):
    """'Qwen Vendor (raw)' -> 'qwen_vendor_raw'. Replicated here (not
    imported) so this file never has to modify invoice_reader.py; this is
    the EXACT same rule used there and in invoice_recorder's header_to_key,
    so all three agree on column names."""
    cleaned = "".join(c if (c.isalnum() or c.isspace()) else " "
                      for c in str(header))
    return "_".join(cleaned.lower().split())


def _load_alias_rows(aliases_path):
    """Read masters\\aliases.csv (or wherever --aliases points) into a plain
    list of dict rows, keeping every column and the original column order.
    A missing file is a normal starting state -- not an error -- and yields
    the default header + an empty list, same spirit as vendor_lookup's
    loaders."""
    fieldnames = list(ALIAS_FIELDNAMES)
    rows = []
    if os.path.exists(aliases_path):
        with open(aliases_path, newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames:
                fieldnames = list(reader.fieldnames)
            for row in reader:
                rows.append(dict(row))
    return fieldnames, rows


def _write_alias_rows(aliases_path, fieldnames, rows):
    """Write the alias file back out safely: build the new content in a
    temp file in the SAME folder, then atomically swap it into place with
    os.replace(). If anything goes wrong while writing, the real file is
    left untouched -- we never leave a half-written aliases.csv behind."""
    directory = os.path.dirname(os.path.abspath(aliases_path)) or "."
    os.makedirs(directory, exist_ok=True)
    fd, tmp_path = tempfile.mkstemp(prefix=".aliases_tmp_", dir=directory)
    try:
        with os.fdopen(fd, "w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()
            for row in rows:
                writer.writerow({name: row.get(name, "") for name in fieldnames})
        os.replace(tmp_path, aliases_path)
    except Exception:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
        raise


def harvest(xlsx_path, aliases_path, master_rows, apply=False,
            name_col="Vendor Name", corrected_key="vendor",
            raw_key="qwen_vendor_raw"):
    """Scan the Invoices sheet in `xlsx_path` for human corrections and turn
    the ones that point at a KNOWN master vendor into alias rows.

    xlsx_path     -- path to the invoices.xlsx to read (never written to).
    aliases_path  -- path to the aliases.csv to read, and (only if apply=True)
                     to update.
    master_rows   -- list of dict rows, straight from vendor_lookup.load_master()
                     (i.e. the caller loads the master; this function does not).
    apply         -- False (default): PREVIEW ONLY, nothing is written.
                     True: masters\\aliases.csv (or --aliases path) is updated.

    Returns a summary dict:
      {
        "corrections_found": int,   # rows where Vendor != Qwen Vendor (raw)
        "aliases_added": int,       # brand-new alias rows (this run)
        "aliases_refreshed": int,   # existing alias rows bumped (this run)
        "new_vendors": [str, ...],  # corrected names NOT in the master --
                                     # reported for a human to review, never
                                     # auto-added anywhere
        "applied": bool,            # whether the aliases file was actually written
        "error": str or None,       # set if the sheet is missing required columns
      }
    """
    today = datetime.now().strftime("%Y-%m-%d")
    master_set = {_normalise(row.get(name_col))
                  for row in master_rows if row.get(name_col)}

    summary = {
        "corrections_found": 0,
        "aliases_added": 0,
        "aliases_refreshed": 0,
        "new_vendors": [],
        "applied": False,
        "error": None,
    }

    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        sheet = (workbook["Invoices"] if "Invoices" in workbook.sheetnames
                 else workbook.active)
        rows_iter = sheet.iter_rows(values_only=True)
        header = next(rows_iter, None) or ()
        columns = {}
        for idx, heading in enumerate(header):
            if heading:
                columns[_key(heading)] = idx

        raw_idx = columns.get(raw_key)
        vendor_idx = columns.get(corrected_key)
        if raw_idx is None or vendor_idx is None:
            summary["error"] = (
                f'Could not find "{corrected_key}" and/or "{raw_key}" columns '
                "in the Invoices sheet -- nothing harvested.")
            return summary

        fieldnames, alias_rows = _load_alias_rows(aliases_path)
        alias_index = {_normalise(row.get("raw", "")): row
                       for row in alias_rows if row.get("raw")}

        seen_new_vendor_keys = set()

        for row in rows_iter:
            raw = row[raw_idx] if raw_idx < len(row) else None
            corrected = row[vendor_idx] if vendor_idx < len(row) else None
            if raw in (None, "") or corrected in (None, ""):
                continue
            raw = str(raw).strip()
            corrected = str(corrected).strip()
            if not raw or not corrected:
                continue
            if _normalise(corrected) == _normalise(raw):
                continue  # no correction actually happened

            summary["corrections_found"] += 1
            normalised_corrected = _normalise(corrected)

            if normalised_corrected in master_set:
                key = _normalise(raw)
                existing = alias_index.get(key)
                if existing is not None:
                    existing["count"] = str(
                        int(existing.get("count") or 0) + 1)
                    existing["last_seen"] = today
                    existing["correct"] = corrected
                    summary["aliases_refreshed"] += 1
                else:
                    new_row = {"raw": raw, "correct": corrected,
                               "count": "1", "last_seen": today}
                    alias_rows.append(new_row)
                    alias_index[key] = new_row
                    summary["aliases_added"] += 1
            else:
                # Rule D5: never auto-alias, never auto-add to the master.
                # Just surface it once for a human to decide on.
                if normalised_corrected not in seen_new_vendor_keys:
                    seen_new_vendor_keys.add(normalised_corrected)
                    summary["new_vendors"].append(corrected)

        if apply:
            _write_alias_rows(aliases_path, fieldnames, alias_rows)
            summary["applied"] = True
    finally:
        workbook.close()

    return summary


def _print_summary(summary, aliases_path, apply):
    print()
    if summary["error"]:
        print(f"STOPPED: {summary['error']}")
        return

    print(f"Corrections found:      {summary['corrections_found']}")
    print(f"Aliases added (new):    {summary['aliases_added']}")
    print(f"Aliases refreshed:      {summary['aliases_refreshed']}")

    if summary["new_vendors"]:
        print(f"\nBrand-new names seen (NOT in the master, NO alias made "
              "-- needs a human OK before adding to the master):")
        for name in summary["new_vendors"]:
            print(f"  - {name}")
    else:
        print("\nNo brand-new names seen.")

    print()
    if apply:
        print(f"APPLIED: {aliases_path} has been updated.")
    else:
        print("PREVIEW ONLY -- nothing was written. Re-run with --apply to "
              f"update {aliases_path}.")


def main():
    parser = argparse.ArgumentParser(
        description="Turn human corrections in invoices.xlsx into "
                     "masters\\aliases.csv entries.")
    parser.add_argument("--xlsx", required=True, type=Path,
                        help="invoices.xlsx to harvest corrections from.")
    parser.add_argument("--aliases", type=Path, default=ALIASES_CSV,
                        help="aliases.csv to read/update. Default: "
                             r"masters\aliases.csv")
    parser.add_argument("--master", type=Path, default=VENDORS_CSV,
                        help="Vendor master to check corrections against. "
                             r"Default: masters\vendors.csv")
    parser.add_argument("--apply", action="store_true",
                        help="Actually update the aliases file. Without this "
                             "flag, PREVIEW ONLY -- nothing is written.")
    parser.add_argument("--property", action="store_true",
                        help="Harvest PROPERTY corrections (Property vs "
                             "Qwen Property (raw)) against masters\\properties.csv "
                             "into masters\\property_aliases.csv, instead of "
                             "vendors.")
    args = parser.parse_args()

    if not args.xlsx.exists():
        sys.exit(f"STOPPED: no spreadsheet at {args.xlsx}")

    if args.property:
        if args.master == VENDORS_CSV:
            args.master = MASTERS_DIR / "properties.csv"
        if args.aliases == ALIASES_CSV:
            args.aliases = MASTERS_DIR / "property_aliases.csv"

    name_col = "Property Name" if args.property else "Vendor Name"
    corrected_key = "property" if args.property else "vendor"
    raw_key = "qwen_property_raw" if args.property else "qwen_vendor_raw"

    master_rows = load_master(str(args.master), name_col=name_col)
    if not master_rows:
        print(f"\n*** WARNING: could not read the master at "
              f"{args.master} (missing or empty). Every correction will be "
              "treated as brand-new since nothing can be confirmed against the "
              "master. ***\n")

    summary = harvest(args.xlsx, args.aliases, master_rows, apply=args.apply,
                      name_col=name_col, corrected_key=corrected_key,
                      raw_key=raw_key)
    _print_summary(summary, args.aliases, args.apply)


if __name__ == "__main__":
    main()
