"""Offline tests for verify_master.py -- the vendor master PROBER.

No browser opens and Bottomline is never contacted. Everything here is plain
data and duck-typed fakes, exactly as selftest_recorder_property.py already
does, so this suite runs on any PC with no login and nothing installed beyond
what the project already uses.

Run:  python prober\\selftest_verify_master.py
"""
import ast
import csv
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent
REPO_ROOT = PROJECT.parent

# TWO folders go on the search path, not one (changed 2026-09-10 when the
# prober moved into its own folder):
#   HERE                  -> prober\, so "import verify_master" works
#   PROJECT / automation  -> automation\, so the shared modules verify_master
#                            imports (fill_plan, harvest_aliases,
#                            invoice_recorder) are findable from here too.
# Dropping the second line makes every test in this file fail at import time
# with ModuleNotFoundError, before a single assertion runs.
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(PROJECT / "automation"))


def run_popup_js(js, heading, window_text, address_block=None,
                 body_inner=None, header_inner=None):
    """Run _POPUP_TEXT_JS in node against a stubbed popup, return what it gives.

    The stub answers querySelector by looking at the selector string, which is
    what lets it model the trap this function keeps falling into: ExtJS gives
    the window HEADER a data-ref="innerCt" as well as the body, and the header
    comes first. Selector MATCHING against real markup is the one thing only a
    browser can settle; every decision this function makes is settled here.

    window_text    innerText of the whole window (heading included)
    address_block  innerText of the element holding the .addressItem lines
    body_inner     innerText of the body's own data-ref="innerCt"
    header_inner   innerText of the HEADER's data-ref="innerCt" -- the decoy
    """
    def node_el(text):
        return "null" if text is None else "{ innerText: %s }" % json.dumps(text)

    folder = Path(tempfile.mkdtemp())
    try:
        script = folder / "run.js"
        script.write_text(
            "const fn = (" + js + ");\n"
            "const addressParent = " + node_el(address_block) + ";\n"
            "const bodyInner = " + node_el(body_inner) + ";\n"
            "const headerInner = " + node_el(header_inner) + ";\n"
            "const addressItem = addressParent "
            "  ? { parentElement: addressParent } : null;\n"
            "const pick = (sel) => {\n"
            "  if (sel.includes('.addressItem')) return addressItem;\n"
            "  if (sel.includes('window_panel') || sel.includes('x-window-body'))\n"
            "    return bodyInner;\n"
            "  if (sel.includes('innerCt')) return headerInner;\n"
            "  return null;\n"
            "};\n"
            "const win = { innerText: " + json.dumps(window_text) + ",\n"
            "  getBoundingClientRect: () => ({ width: 400, height: 156 }),\n"
            "  querySelector: pick };\n"
            "global.document = { querySelector: () => win };\n"
            "process.stdout.write(JSON.stringify(fn(["
            + json.dumps("[auto_id=\"vendor_address_window\"]") + ", "
            + json.dumps(heading) + "])));\n",
            encoding="utf-8")
        done = subprocess.run(["node", str(script)],
                              capture_output=True, text=True)
        if done.returncode != 0:
            raise AssertionError(f"node could not run the popup JS:\n{done.stderr}")
        return json.loads(done.stdout)
    finally:
        shutil.rmtree(folder, ignore_errors=True)


def js_blocks():
    """Every _JS constant in verify_master.py, as it is WRITTEN ON DISK.

    The source text, not the parsed value, because what is being checked is
    the quoting itself -- whether the literal is a raw string. Reading the
    parsed value would hide exactly the bug this exists to catch.
    """
    source = (HERE / "verify_master.py").read_text(encoding="utf-8")
    tree = ast.parse(source)
    blocks = {}
    for node in ast.walk(tree):
        if isinstance(node, ast.Assign):
            for target in node.targets:
                name = getattr(target, "id", "")
                if name.endswith("_JS"):
                    blocks[name] = ast.get_source_segment(source, node.value)
    return blocks

from harvest_aliases import normalise_package as normalise_master_package  # noqa: E402
from verify_master import (                                                # noqa: E402
    AMBIGUOUS,
    BATCH_SIZE,
    CODE_MISMATCH,
    CONFIRMED,
    FIELD_LABELS,
    NOT_FOUND,
    NO_CODE,
    PACKAGE_LEARNED,
    PAUSE_QUIT,
    PAUSE_RETRY,
    PAUSE_SKIP,
    PAUSE_STOP_ASKING,
    POPUP_HEADING,
    PROBED_FIELDNAMES,
    REVIEW_HEADERS,
    REVIEW_SHEETS,
    UNVERIFIED_PACKAGE,
    VERDICT_MARK,
    WRONG_PACKAGE,
    append_probed_rows,
    archive_probe_csv,
    attempt,
    calibration_result,
    calibration_verdict,
    classify_probe,
    clear_code_box,
    codes_match,
    form_confirms,
    gather_visible_data,
    describe_probe,
    indexing_form_is_open,
    is_phone_line,
    load_master_rows,
    make_batches,
    normalise_code,
    normalise_name_for_compare,
    package_already_selected,
    parse_vendor_popup,
    pause_message,
    pause_until_loaded,
    pick_calibration_rows,
    probed_row,
    read_box_value,
    read_checkpoint,
    read_pause_choice,
    resume_start_index,
    review_sheet_for,
    route_to_sheets,
    screen_vendor_name,
    rows_in_master_order,
    summary_lines,
    write_checkpoint,
    write_review_workbook,
)


def vendor(name, code="", package="", status="ACTIVE"):
    """One master row, shaped exactly like a csv.DictReader row from
    masters/vendors.csv -- same nine columns, same spellings."""
    return {"Vendor Name": name, "City": "", "State": "", "Zip": "",
            "Membership ID": "", "ERP Vendor ID": code, "Address": "",
            "Integration Package": package, "Status": status}


def hit(package="YARDI", code="V1", name="ACME LLC", city="", state="", postcode=""):
    """A probe where the popup showed back the code we typed."""
    return attempt(package, code, 1,
                   {"name": name, "code": code, "street": "",
                    "city": city, "state": state, "zip": postcode})


def miss(package="YARDI", code="V1", count=0):
    """A probe where no popup was read at all."""
    return attempt(package, code, count, {})


def other_vendor(package="YARDI", code="V1", popup_code="V999", name="SOMEONE ELSE"):
    """A probe where a popup opened but named a DIFFERENT code."""
    return attempt(package, code, 1, {"name": name, "code": popup_code})


# ==========================================================================
# Task 1 -- the guard that has to exist before the prober can write anything
# ==========================================================================

class OutputsAreGitignored(unittest.TestCase):
    """Hard rule #1. vendors_probed.csv and master_review.xlsx hold REAL vendor
    names and now real addresses too, so they must be ignored BEFORE the prober
    can ever write them. The existing '**/masters/alias_review.xlsx' line does
    not cover them -- it names one file, not a pattern.
    """

    def setUp(self):
        self.text = (REPO_ROOT / ".gitignore").read_text(encoding="utf-8")

    def test_probed_csv_is_ignored(self):
        self.assertIn("**/masters/vendors_probed.csv", self.text)

    def test_review_workbook_is_ignored(self):
        self.assertIn("**/masters/master_review.xlsx", self.text)

    def test_checkpoint_is_ignored(self):
        self.assertIn("**/masters/.probe_checkpoint.json", self.text)

    def test_the_archive_folder_is_ignored(self):
        """Added 2026-09-09 with archive_probe_csv(). An archived run is a
        byte-for-byte copy of vendors_probed.csv, so it carries exactly the same
        real vendor names and addresses. Ignoring the live file but not its
        archive would leak the identical data through a different path."""
        self.assertIn("**/masters/archive/", self.text)


# ==========================================================================
# Task 2 -- reading, sorting and batching the master
# ==========================================================================

class RowsInMasterOrder(unittest.TestCase):
    """Changed 2026-09-09 on the operator's instruction. Rows are probed in the
    master's OWN order -- row 1 of a run is row 1 of vendors.csv, batch 1 is
    rows 1-25 -- so the operator can follow a live run line by line in the
    spreadsheet.

    The sort this replaced grouped the packages together, which cut the run
    from ~959 package switches to 3. That saving is genuinely gone, and these
    tests exist so nobody quietly puts it back: a sorted run is faster and
    impossible to check against the file.
    """

    def test_the_masters_order_is_kept_exactly(self):
        rows = [vendor("A", package="INTACCT (1)"), vendor("B", package="YARDI (2)"),
                vendor("C", package=""), vendor("D", package="INTACCT (1)"),
                vendor("E", package="YARDI (2)")]
        self.assertEqual([r["Vendor Name"] for r in rows_in_master_order(rows)],
                         ["A", "B", "C", "D", "E"])

    def test_the_package_never_influences_the_order(self):
        """The exact case the old sort reordered: a no-package row sitting
        between two YARDI rows used to be moved to the very end."""
        rows = [vendor("A", package="YARDI (2)"), vendor("B", package=""),
                vendor("C", package="YARDI (2)")]
        self.assertEqual([r["Vendor Name"] for r in rows_in_master_order(rows)],
                         ["A", "B", "C"])

    def test_the_first_batch_is_the_first_25_rows_of_the_file(self):
        """What the operator actually relies on while a run is going."""
        rows = [vendor(f"R{n}") for n in range(60)]
        first = make_batches(rows_in_master_order(rows), 25)[0]
        self.assertEqual([r["Vendor Name"] for r in first],
                         [f"R{n}" for n in range(25)])

    def test_the_callers_list_is_not_handed_straight_back(self):
        """A copy, so appending to the result cannot mutate what
        load_master_rows() returned."""
        rows = [vendor("A")]
        self.assertIsNot(rows_in_master_order(rows), rows)

    def test_no_rows_gives_no_rows(self):
        self.assertEqual(rows_in_master_order([]), [])


class MakeBatches(unittest.TestCase):
    def test_959_rows_at_25_gives_39_batches(self):
        self.assertEqual(len(make_batches([vendor(str(n)) for n in range(959)], 25)),
                         39)

    def test_the_last_batch_is_short(self):
        batches = make_batches([vendor(str(n)) for n in range(959)], 25)
        self.assertEqual(len(batches[-1]), 9)
        self.assertTrue(all(len(b) == 25 for b in batches[:-1]))

    def test_an_exact_multiple_leaves_no_empty_batch(self):
        self.assertEqual([len(b) for b in make_batches(
            [vendor(str(n)) for n in range(50)], 25)], [25, 25])

    def test_no_rows_gives_no_batches(self):
        self.assertEqual(make_batches([], 25), [])

    def test_a_batch_size_below_one_is_refused(self):
        with self.assertRaises(ValueError):
            make_batches([vendor("A")], 0)

    def test_default_batch_size_is_25(self):
        self.assertEqual(BATCH_SIZE, 25)


class LoadMasterRows(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.mkdtemp()
        self.path = Path(self.folder) / "vendors.csv"

    def tearDown(self):
        shutil.rmtree(self.folder, ignore_errors=True)

    def test_reads_the_real_columns(self):
        with self.path.open("w", newline="", encoding="utf-8") as out:
            writer = csv.writer(out)
            writer.writerow(["Vendor Name", "City", "State", "Zip",
                             "Membership ID", "ERP Vendor ID", "Address",
                             "Integration Package", "Status"])
            writer.writerow(["ACME LLC", "", "", "", "", "V0012", "",
                             "YARDI (2)", "ACTIVE"])
        rows = load_master_rows(self.path)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["ERP Vendor ID"], "V0012")
        self.assertEqual(rows[0]["Integration Package"], "YARDI (2)")

    def test_a_missing_file_gives_no_rows(self):
        self.assertEqual(load_master_rows(self.path / "nope.csv"), [])


# ==========================================================================
# The Vendor Address popup -- the witness everything else now rests on
# ==========================================================================

REAL_POPUP = "\n".join([
    "Vendor Address",
    "American Pest Management (4339-001)",
    "6940 Columbia Gateway Drive",
    "#250",
    "Columbia, MD 21046",
    "(301) 891-2600",
])


class ParseVendorPopup(unittest.TestCase):
    """Fed the exact text captured from the live site on 2026-09-09, and the
    same format the Master Maker handover recorded (HANDOFF.md:90).
    """

    def setUp(self):
        self.found = parse_vendor_popup(REAL_POPUP)

    def test_the_name_is_split_from_the_code(self):
        self.assertEqual(self.found["name"], "American Pest Management")
        self.assertEqual(self.found["code"], "4339-001")

    def test_the_city_state_and_zip_come_out_separately(self):
        self.assertEqual(self.found["city"], "Columbia")
        self.assertEqual(self.found["state"], "MD")
        self.assertEqual(self.found["zip"], "21046")

    def test_the_street_keeps_both_of_its_lines(self):
        self.assertEqual(self.found["street"], "6940 Columbia Gateway Drive #250")

    def test_the_phone_number_is_never_collected(self):
        """The user asked for the phone to be skipped, and it is client contact
        data we have no reason to hold."""
        self.assertNotIn("301", self.found["raw"])
        self.assertNotIn("891", str(self.found))

    def test_the_heading_line_is_dropped(self):
        self.assertFalse(self.found["raw"].lower().startswith(
            POPUP_HEADING.lower()))

    def test_a_close_glyph_is_dropped(self):
        found = parse_vendor_popup("Vendor Address\n✕\nACME LLC (V1)")
        self.assertEqual(found["name"], "ACME LLC")
        self.assertEqual(found["code"], "V1")

    def test_empty_text_gives_empty_fields_not_a_crash(self):
        found = parse_vendor_popup("")
        self.assertEqual(found["name"], "")
        self.assertEqual(found["code"], "")

    def test_none_gives_empty_fields(self):
        self.assertEqual(parse_vendor_popup(None)["code"], "")

    def test_a_name_with_no_bracketed_code_still_yields_the_name(self):
        found = parse_vendor_popup("Vendor Address\nACME LLC")
        self.assertEqual(found["name"], "ACME LLC")
        self.assertEqual(found["code"], "")

    def test_a_name_that_itself_contains_brackets(self):
        found = parse_vendor_popup(
            "Vendor Address\nACME (HOLDINGS) LLC (V1)\nColumbia, MD 21046")
        self.assertEqual(found["name"], "ACME (HOLDINGS) LLC")
        self.assertEqual(found["code"], "V1")

    def test_a_zip_plus_four_is_read(self):
        found = parse_vendor_popup(
            "Vendor Address\nACME (V1)\nColumbia, MD 21046-1234")
        self.assertEqual(found["zip"], "21046-1234")

    def test_a_single_line_popup_is_not_mistaken_for_a_city(self):
        found = parse_vendor_popup("Vendor Address\nACME LLC (V1)")
        self.assertEqual(found["name"], "ACME LLC")
        self.assertEqual(found["city"], "")


class TheButtonIsNotMistakenForThePopup(unittest.TestCase):
    """THE BUG THAT KILLED THE FIRST RUN TO REACH A REAL VENDOR, 2026-09-09.

    POPUP_HEADING and POPUP_BUTTON are the same words -- "Vendor Address" --
    and the control that opens the popup is a div, not a real button
    (content.js:49), so it is inside the very querySelectorAll that hunts for
    the popup. Its text is 14 characters; the popup's is about 100. The finder
    kept the SHORTEST match, so the button won every time -- even with the
    popup open and correct on screen. parse_vendor_popup() then stripped
    "Vendor Address" as the heading, nothing was left, and a vendor that had
    loaded perfectly was reported as "showed nothing we could read".

    It also made the 2.5-second wait meaningless: the button is on screen
    instantly, so the polling loop broke on its first pass and never waited for
    the popup -- the "opens and closes too fast to capture" the operator
    watched happen.
    """

    def setUp(self):
        self.blocks = js_blocks()

    def test_the_heading_and_the_button_really_are_the_same_words(self):
        """If these ever stop matching, the bug below stops being possible and
        the guard could be relaxed -- so the assumption is written down."""
        from verify_master import POPUP_BUTTON
        self.assertEqual(POPUP_HEADING, POPUP_BUTTON)

    def test_every_js_block_is_a_raw_string(self):
        r"""THE SECOND BUG, AND THE WORSE ONE -- found live 2026-09-10.

        The old _POPUP_TEXT_JS was a normal Python string, and the comment
        inside it warning about \n contained a real \n. Python turned it into
        a line break, the comment broke in half, and what was left of it --
        "': this whole block is a PYTHON string" -- became a stray quote in
        the middle of the code. The browser answered with

            SyntaxError: Invalid or unexpected token

        every single time, so the JavaScript NEVER RAN. popup_text() caught
        that exception, returned "", and told nobody. Every vendor was
        reported as "showed nothing we could read", and close_vendor_popup()
        -- which asked popup_text() whether a popup existed -- concluded there
        was nothing to close and never clicked the x. Three days of symptoms,
        one broken comment.

        A raw string cannot do that: backslashes reach the browser as written.
        """
        for name, source in self.blocks.items():
            with self.subTest(block=name):
                self.assertTrue(
                    source.startswith('r"""') or source.startswith("r'''"),
                    f"{name} must be a raw string (r\"\"\") so a backslash "
                    f"cannot be re-interpreted on the way to the browser")

    @unittest.skipUnless(shutil.which("node"), "node is not installed")
    def test_a_heading_with_no_body_under_it_is_not_an_answer(self):
        """THE 2026-09-10 REGRESSION, RUN FOR REAL IN NODE.

        ExtJS paints the window frame first and fills the body in when
        getVendor returns. For the first few hundred milliseconds the whole
        window reads exactly "Vendor Address". The live run read that at
        240ms, the polling loop took it as success and stopped waiting,
        parse_vendor_popup() stripped it as the heading, every field came
        back empty and the popup was closed -- it "opened and shut in a
        fraction of a second".

        So: heading alone must return "" (keep waiting), and a real body must
        come back whole. This runs the ACTUAL JavaScript, not a description
        of it.
        """
        from verify_master import POPUP_HEADING, _POPUP_TEXT_JS

        # 1. The frame is up, nothing in it yet, and the HEADER's innerCt is
        #    sitting there reading "Vendor Address" -- the decoy that broke
        #    the 2026-09-10 run. Must be refused so the caller keeps waiting.
        frame_only = run_popup_js(
            _POPUP_TEXT_JS, POPUP_HEADING,
            window_text="Vendor Address", header_inner="Vendor Address")
        self.assertEqual(frame_only, "",
                         "the window header must never be read as the answer")

        # 2. Loaded, real markup (operator's capture, 2026-09-10). The header
        #    decoy is STILL present -- the address block must win anyway.
        loaded = run_popup_js(
            _POPUP_TEXT_JS, POPUP_HEADING,
            window_text="Vendor Address\nA/C Control Inc (3983-001)",
            header_inner="Vendor Address",
            address_block=("A/C Control Inc (3983-001)\n"
                           "6649 Garland Drive, Suite 9\n"
                           "Warrenton, VA 20187"))
        found = parse_vendor_popup(loaded)
        self.assertEqual(found["code"], "3983-001")
        self.assertEqual(found["name"], "A/C Control Inc")
        self.assertEqual(found["street"], "6649 Garland Drive, Suite 9")
        self.assertEqual(found["city"], "Warrenton")
        self.assertEqual(found["state"], "VA")
        self.assertEqual(found["zip"], "20187")

        # 3. No .addressItem markup, but the body's own content area has the
        #    lines -- the next candidate down must pick them up.
        body_only = run_popup_js(
            _POPUP_TEXT_JS, POPUP_HEADING,
            window_text="Vendor Address\nIllustratus (762-001)",
            header_inner="Vendor Address",
            body_inner="Illustratus (762-001)\n8455 Lenexa Drive\nLenexa, KS 66214")
        self.assertEqual(parse_vendor_popup(body_only)["code"], "762-001")

        # 4. Nothing scoped matches at all, but the window itself carries a
        #    real body: the last-resort candidate answers, heading included
        #    for parse_vendor_popup() to strip.
        fallback = run_popup_js(
            _POPUP_TEXT_JS, POPUP_HEADING,
            window_text="Vendor Address\nIllustratus (762-001)\n8455 Lenexa Drive")
        self.assertEqual(parse_vendor_popup(fallback)["name"], "Illustratus")

    @unittest.skipUnless(shutil.which("node"), "node is not installed")
    def test_every_js_block_actually_parses_as_javascript(self):
        """The test that would have caught the above on the day it was written.

        Every _JS constant is handed to node --check exactly as Python would
        hand it to the browser. No browser, no login, no Bottomline -- it runs
        anywhere node is installed, and is skipped where it is not.
        """
        from verify_master import (_CLOSE_POPUP_JS, _FIELD_BY_LABEL_JS,
                                   _FIND_BUTTON_JS, _POPUP_OPEN_JS,
                                   _POPUP_TEXT_JS)
        blocks = {"_FIELD_BY_LABEL_JS": _FIELD_BY_LABEL_JS,
                  "_POPUP_TEXT_JS": _POPUP_TEXT_JS,
                  "_POPUP_OPEN_JS": _POPUP_OPEN_JS,
                  "_CLOSE_POPUP_JS": _CLOSE_POPUP_JS,
                  "_FIND_BUTTON_JS": _FIND_BUTTON_JS}
        for name, js in blocks.items():
            with self.subTest(block=name):
                folder = Path(tempfile.mkdtemp())
                try:
                    script = folder / "block.js"
                    # Wrapped the way page.evaluate() wraps it: an expression.
                    script.write_text(f"const fn = ({js});\n", encoding="utf-8")
                    done = subprocess.run(["node", "--check", str(script)],
                                          capture_output=True, text=True)
                    self.assertEqual(done.returncode, 0,
                                     f"{name} is not valid JavaScript:\n"
                                     f"{done.stderr}")
                finally:
                    shutil.rmtree(folder, ignore_errors=True)

    def test_a_bare_label_parses_to_nothing_which_is_what_went_wrong(self):
        """Kept as the record of the symptom: this is exactly what
        parse_vendor_popup() was being handed on every single row."""
        found = parse_vendor_popup("Vendor Address")
        self.assertEqual(found["name"], "")
        self.assertEqual(found["code"], "")

    def test_a_real_popup_still_parses_as_it_always_did(self):
        found = parse_vendor_popup(REAL_POPUP)
        self.assertEqual(found["code"], "4339-001")


SCREEN = {"integration_package": "INTACCT", "property_code": "0000",
          "property": "000 - PROPERTY NOT FOUND",
          "vendor_code": "762-001", "vendor": "Illustratus (Lenexa/66214)"}


class ReadingTheFormByLabel(unittest.TestCase):
    """The Master Maker extension's method, adopted 2026-09-09 on the operator's
    instruction.

    The extension NEVER opens or closes a popup (content.js:207) -- a person
    clicks and it reads what is on screen. Everything but the addresses it takes
    straight off the form by label (content.js:122-131). The form already shows
    what we were opening a modal to learn:

        Vendor Code : 762-001
        Vendor      : Illustratus (Lenexa/66214)

    So the witness moved to the form, and the popup now supplies only the street
    line the operator needs for the master.
    """

    def test_the_five_labels_are_the_extensions_own(self):
        self.assertEqual([label for _, label in FIELD_LABELS],
                         ["Integration Package", "Property Code", "Property",
                          "Vendor Code", "Vendor"])

    def test_the_form_confirms_a_vendor_that_really_loaded(self):
        self.assertTrue(form_confirms(attempt("INTACCT", "762-001", 1,
                                              form=SCREEN)))

    def test_a_code_box_still_holding_the_PREVIOUS_code_is_not_a_hit(self):
        """The failure the popup was originally introduced to catch. If the box
        did not clear it echoes the OLD code, which is not the one we typed."""
        self.assertFalse(form_confirms(attempt("INTACCT", "999-999", 1,
                                               form=SCREEN)))

    def test_a_code_that_resolved_to_nothing_is_not_a_hit(self):
        """The vendor box stays empty, which is the extension's own
        completeness test at content.js:98."""
        blank = dict(SCREEN, vendor="")
        self.assertFalse(form_confirms(attempt("INTACCT", "762-001", 1,
                                               form=blank)))

    def test_the_form_alone_counts_as_a_hit_with_no_popup_at_all(self):
        """The point of the rework: a jammed or unread popup can no longer turn
        a vendor that plainly loaded into a NOT_FOUND."""
        one = attempt("INTACCT", "762-001", 1, form=SCREEN)
        self.assertEqual(one["popup_code"], "")
        self.assertEqual(
            classify_probe("762-001", "INTACCT (1)", [one])[0], CONFIRMED)

    def test_the_popup_is_still_a_valid_witness_on_its_own(self):
        """Rows probed before the rework, and any row where the form cannot be
        read, still resolve on the popup exactly as they always did."""
        self.assertEqual(
            classify_probe("V1", "YARDI (2)", [hit("YARDI", "V1", "ACME")])[0],
            CONFIRMED)

    def test_the_five_fields_are_carried_on_the_attempt(self):
        one = attempt("INTACCT", "762-001", 1, form=SCREEN)
        self.assertEqual(one["form_vendor"], "Illustratus (Lenexa/66214)")
        self.assertEqual(one["form_property_code"], "0000")
        self.assertEqual(one["form_package"], "INTACCT")

    def test_missing_form_data_defaults_to_empty_not_absent(self):
        self.assertEqual(attempt("YARDI", "V1", 1)["form_vendor"], "")

    def test_the_form_is_read_BEFORE_any_popup_is_opened(self):
        """Source-level, because the ORDER is the whole safety property: if the
        popup jams afterwards, the row has already been captured."""
        source = (HERE / "verify_master.py").read_text(encoding="utf-8")
        body = source[source.index("def probe_one_code("):]
        body = body[:body.index("\n# ---")]
        self.assertLess(body.index("gather_visible_data(page)"),
                        body.index("open_and_read_vendor_popup("))

    def test_the_label_matching_is_strict_so_Vendor_is_not_Vendor_Code(self):
        """content.js:145. A 'contains' match would read the Vendor Code box
        for the label 'Vendor' and every row would look self-consistent."""
        source = (HERE / "verify_master.py").read_text(encoding="utf-8")
        self.assertIn("clean(node.textContent) === cleanTarget", source)


class ScreenVendorName(unittest.TestCase):
    """The Vendor box reads 'Name (City/ZIP)'. The brackets hold a PLACE, not a
    code -- unlike the popup, where they hold '(4339-001)'. Keeping the two
    apart is what stops a city being recorded as a vendor code."""

    def test_the_name_is_taken_from_before_the_brackets(self):
        self.assertEqual(screen_vendor_name("Illustratus (Lenexa/66214)"),
                         "Illustratus")

    def test_a_name_without_brackets_survives_unchanged(self):
        self.assertEqual(screen_vendor_name("ACME LLC"), "ACME LLC")

    def test_blank_and_none_are_safe(self):
        self.assertEqual(screen_vendor_name(""), "")
        self.assertEqual(screen_vendor_name(None), "")

    def test_it_lets_the_spelling_check_work_with_no_popup(self):
        """--no-address must not silently switch the Name Differs sheet off."""
        row = probed("Illustratus Inc", CONFIRMED, "", code="762-001",
                     popup_code="762-001", screen="Illustratus (Lenexa/66214)")
        self.assertEqual(row["Bottomline Name"], "Illustratus")
        self.assertEqual(review_sheet_for(row), "Name Differs")

    def test_the_popup_name_still_wins_when_there_is_one(self):
        row = probed("ACME", CONFIRMED, "ACME LIMITED", code="V1",
                     screen="SOMETHING ELSE (Lenexa/66214)")
        self.assertEqual(row["Bottomline Name"], "ACME LIMITED")


class IsPhoneLine(unittest.TestCase):
    """Kept strict on purpose: a street number or a ZIP must never be mistaken
    for a phone number and dropped."""

    def test_a_bracketed_phone_is_a_phone(self):
        self.assertTrue(is_phone_line("(301) 891-2600"))

    def test_a_plain_dashed_phone_is_a_phone(self):
        self.assertTrue(is_phone_line("301-891-2600"))

    def test_a_city_line_is_not_a_phone(self):
        self.assertFalse(is_phone_line("Columbia, MD 21046"))

    def test_a_street_line_is_not_a_phone(self):
        self.assertFalse(is_phone_line("6940 Columbia Gateway Drive"))

    def test_a_zip_alone_is_too_short_to_be_a_phone(self):
        self.assertFalse(is_phone_line("21046"))

    def test_a_suite_number_is_not_a_phone(self):
        self.assertFalse(is_phone_line("#250"))


class CodesMatch(unittest.TestCase):
    def test_case_and_spacing_are_forgiven(self):
        self.assertTrue(codes_match(" v1 ", "V1"))

    def test_hyphens_are_significant(self):
        """'4339-001' and '4339001' are different codes."""
        self.assertFalse(codes_match("4339-001", "4339001"))

    def test_a_blank_never_matches_anything(self):
        self.assertFalse(codes_match("", ""))
        self.assertFalse(codes_match("", "V1"))

    def test_normalise_code_handles_none(self):
        self.assertEqual(normalise_code(None), "")


# ==========================================================================
# Turning one probe into a verdict
# ==========================================================================

class ClassifyProbe(unittest.TestCase):
    """One rule runs through all of it: a match counts only when the popup
    showed back THE CODE WE TYPED.

    Not "a popup appeared" -- Bottomline opens the Vendor Address popup even on
    an incomplete form (Master Maker HANDOFF.md:79), so a dead code still pops
    a dialog. Not "the name box is filled" either -- the code box does not
    always clear, so that box can still hold the previous vendor.
    """

    def test_the_popup_showing_our_code_under_the_claimed_package(self):
        verdict, package, best = classify_probe(
            "V1", "YARDI (2)", [hit("YARDI", "V1", "ACME LLC")])
        self.assertEqual(verdict, CONFIRMED)
        self.assertEqual(package, "YARDI")
        self.assertEqual(best["popup_name"], "ACME LLC")

    def test_nothing_then_our_code_across_the_packages_is_wrong_package(self):
        verdict, package, best = classify_probe(
            "V1", "YARDI (2)",
            [miss("YARDI", "V1"), hit("INTACCT", "V1", "ACME LLC")])
        self.assertEqual(verdict, WRONG_PACKAGE)
        self.assertEqual(package, "INTACCT")

    def test_no_popup_under_either_package_is_not_found(self):
        verdict, package, best = classify_probe(
            "V1", "YARDI (2)", [miss("YARDI", "V1"), miss("INTACCT", "V1")])
        self.assertEqual(verdict, NOT_FOUND)
        self.assertEqual((package, best), ("", {}))

    def test_a_popup_naming_a_different_code_is_a_mismatch_not_not_found(self):
        """THE clearing bug, made visible. If the code box did not clear, the
        popup names the PREVIOUS vendor. Calling that NOT_FOUND would bury it."""
        verdict, package, best = classify_probe(
            "V1", "YARDI (2)", [other_vendor("YARDI", "V1", "V999", "PREVIOUS CO")])
        self.assertEqual(verdict, CODE_MISMATCH)
        self.assertEqual(best["popup_code"], "V999")
        self.assertEqual(best["popup_name"], "PREVIOUS CO")

    def test_a_real_hit_beats_a_mismatch_seen_under_the_other_package(self):
        verdict, _, _ = classify_probe(
            "V1", "YARDI (2)",
            [other_vendor("YARDI", "V1", "V999"), hit("INTACCT", "V1")])
        self.assertEqual(verdict, WRONG_PACKAGE)

    def test_two_or_more_dropdown_matches_is_ambiguous(self):
        verdict, _, _ = classify_probe(
            "V1", "YARDI (2)", [attempt("YARDI", "V1", 3, {"name": "A", "code": "V1"})])
        self.assertEqual(verdict, AMBIGUOUS)

    def test_a_blank_code_is_never_probed(self):
        verdict, package, best = classify_probe("", "YARDI (2)", [])
        self.assertEqual(verdict, NO_CODE)
        self.assertEqual((package, best), ("", {}))

    def test_a_whitespace_only_code_is_also_no_code(self):
        self.assertEqual(classify_probe("   ", "YARDI (2)", [])[0], NO_CODE)

    def test_no_claimed_package_and_one_match_learns_the_package(self):
        verdict, package, _ = classify_probe(
            "V1", "", [miss("YARDI", "V1"), hit("INTACCT", "V1")])
        self.assertEqual(verdict, PACKAGE_LEARNED)
        self.assertEqual(package, "INTACCT")

    def test_no_claimed_package_and_a_match_under_both_teaches_nothing(self):
        verdict, _, _ = classify_probe(
            "V1", "", [hit("YARDI", "V1"), hit("INTACCT", "V1")])
        self.assertEqual(verdict, UNVERIFIED_PACKAGE)

    def test_the_claimed_package_is_matched_through_the_bracketed_number(self):
        """'YARDI (2)' in the master and 'YARDI' on screen are the same package.
        Without this, every single row would come back WRONG_PACKAGE."""
        self.assertEqual(
            classify_probe("V1", "INTACCT (1)", [hit("INTACCT", "V1")])[0],
            CONFIRMED)

    def test_when_the_code_box_does_not_filter_the_package_is_unverified(self):
        verdict, package, best = classify_probe(
            "V1", "YARDI (2)", [hit("YARDI", "V1", "ACME LLC")],
            code_box_filters=False)
        self.assertEqual(verdict, UNVERIFIED_PACKAGE)
        self.assertEqual(package, "")
        self.assertEqual(best["popup_name"], "ACME LLC")

    def test_an_unfiltered_box_still_reports_a_dead_code(self):
        self.assertEqual(
            classify_probe("V1", "YARDI (2)", [miss("YARDI", "V1")],
                           code_box_filters=False)[0], NOT_FOUND)

    def test_an_unfiltered_box_still_reports_a_mismatch(self):
        self.assertEqual(
            classify_probe("V1", "YARDI (2)", [other_vendor("YARDI", "V1")],
                           code_box_filters=False)[0], CODE_MISMATCH)


# ==========================================================================
# Phase 0 -- the calibration that can stop the run
# ==========================================================================

class PickCalibrationRows(unittest.TestCase):
    """Phase 0 is only as good as the 10 rows it runs on. A row with a duplicate
    name, no code, or a non-ACTIVE status adds a second explanation for every
    result, so those rows never enter it.
    """

    def test_takes_five_from_each_package(self):
        rows = ([vendor(f"Y{n}", code=f"Y{n}", package="YARDI (2)") for n in range(9)]
                + [vendor(f"I{n}", code=f"I{n}", package="INTACCT (1)") for n in range(9)])
        packages = [normalise_master_package(r["Integration Package"])
                    for r in pick_calibration_rows(rows)]
        self.assertEqual(packages.count("YARDI"), 5)
        self.assertEqual(packages.count("INTACCT"), 5)

    def test_rows_without_a_code_are_never_picked(self):
        rows = [vendor("A", code="", package="YARDI (2)"),
                vendor("B", code="Y1", package="YARDI (2)")]
        self.assertEqual([r["Vendor Name"] for r in pick_calibration_rows(rows)], ["B"])

    def test_inactive_rows_are_never_picked(self):
        rows = [vendor("A", code="Y1", package="YARDI (2)", status="INACTIVE"),
                vendor("B", code="Y2", package="YARDI (2)")]
        self.assertEqual([r["Vendor Name"] for r in pick_calibration_rows(rows)], ["B"])

    def test_a_name_that_appears_twice_is_never_picked(self):
        rows = [vendor("DUP", code="Y1", package="YARDI (2)"),
                vendor("DUP", code="Y2", package="INTACCT (1)"),
                vendor("SOLO", code="Y3", package="YARDI (2)")]
        self.assertEqual([r["Vendor Name"] for r in pick_calibration_rows(rows)],
                         ["SOLO"])

    def test_no_package_rows_are_never_picked(self):
        rows = [vendor("A", code="X1", package=""),
                vendor("B", code="Y1", package="YARDI (2)")]
        self.assertEqual([r["Vendor Name"] for r in pick_calibration_rows(rows)], ["B"])

    def test_returns_what_it_can_when_a_package_is_short(self):
        self.assertEqual(
            len(pick_calibration_rows([vendor("Y0", code="Y0", package="YARDI (2)")])),
            1)


class CalibrationVerdict(unittest.TestCase):
    """The four-row table in spec section 4. The two STOP rows matter most:
    they are the cases where crawling would produce 959 confident, wrong rows.
    """

    def test_each_code_matches_only_its_claimed_package(self):
        action, filters, reason = calibration_verdict(
            [calibration_result("YARDI", ["YARDI"]),
             calibration_result("INTACCT", ["INTACCT"])])
        self.assertEqual(action, "run")
        self.assertTrue(filters)
        self.assertIn("filters", reason)

    def test_the_claimed_package_is_read_through_the_bracketed_number(self):
        self.assertEqual(calibration_verdict(
            [calibration_result("YARDI (2)", ["YARDI"]),
             calibration_result("INTACCT (1)", ["INTACCT"])])[0], "run")

    def test_matching_under_both_packages_runs_but_unverified(self):
        action, filters, reason = calibration_verdict(
            [calibration_result("YARDI", ["YARDI", "INTACCT"]),
             calibration_result("INTACCT", ["YARDI", "INTACCT"])])
        self.assertEqual(action, "run")
        self.assertFalse(filters)
        self.assertIn("not", reason.lower())

    def test_matching_the_wrong_package_stops_the_run(self):
        action, _, reason = calibration_verdict(
            [calibration_result("YARDI", ["INTACCT"]),
             calibration_result("INTACCT", ["YARDI"])])
        self.assertEqual(action, "stop")
        self.assertIn("package column", reason)

    def test_mixed_results_stop_the_run(self):
        self.assertEqual(calibration_verdict(
            [calibration_result("YARDI", ["YARDI"]),
             calibration_result("INTACCT", ["YARDI", "INTACCT"])])[0], "stop")

    def test_a_code_matching_nothing_stops_the_run(self):
        self.assertEqual(calibration_verdict(
            [calibration_result("YARDI", ["YARDI"]),
             calibration_result("INTACCT", [])])[0], "stop")

    def test_no_results_at_all_stops_the_run(self):
        self.assertEqual(calibration_verdict([])[0], "stop")


# ==========================================================================
# Writing every batch, and resuming
# ==========================================================================

def probed(name, verdict, bottomline_name="", package="YARDI (2)",
           matched="YARDI", code="V1", popup_code=None,
           city="", state="", postcode="", box="", screen=""):
    best = {"popup_name": bottomline_name,
            "popup_code": code if popup_code is None else popup_code,
            "box": box,
            "form_vendor": screen,
            "street": "", "city": city, "state": state, "zip": postcode}
    return probed_row(vendor(name, code=code, package=package),
                      verdict, matched, best, 1, [], "2026-09-09 10:00")


class AppendProbedRows(unittest.TestCase):
    """The whole reason for batching. On 2026-09-07 the reader lost an entire
    run because it saved only at the end. If this ever truncates instead of
    appending, batching buys nothing and that bug is back.
    """

    def setUp(self):
        self.folder = tempfile.mkdtemp()
        self.path = Path(self.folder) / "vendors_probed.csv"

    def tearDown(self):
        shutil.rmtree(self.folder, ignore_errors=True)

    def _rows(self):
        with self.path.open(newline="", encoding="utf-8-sig") as handle:
            return list(csv.DictReader(handle))

    def test_the_first_write_creates_the_file_with_a_header(self):
        append_probed_rows(self.path, [probed("ACME LLC", CONFIRMED, "ACME LLC")])
        self.assertTrue(self.path.exists())
        self.assertEqual(self._rows()[0]["Verdict"], CONFIRMED)

    def test_a_second_batch_is_appended_not_written_over(self):
        for name in ("A", "B"):
            append_probed_rows(self.path, [probed(name, CONFIRMED, name)])
        self.assertEqual([r["Vendor Name"] for r in self._rows()], ["A", "B"])

    def test_the_header_is_written_once_only(self):
        for name in ("A", "B"):
            append_probed_rows(self.path, [probed(name, CONFIRMED, name)])
        self.assertEqual(
            self.path.read_text(encoding="utf-8").count("Vendor Name"), 1)

    def test_writing_no_rows_does_nothing(self):
        append_probed_rows(self.path, [])
        self.assertFalse(self.path.exists())

    def test_a_missing_folder_is_created(self):
        deeper = Path(self.folder) / "masters" / "vendors_probed.csv"
        append_probed_rows(deeper, [probed("A", CONFIRMED, "A")])
        self.assertTrue(deeper.exists())

    def test_every_fieldname_is_present_in_a_built_row(self):
        self.assertEqual(sorted(probed("A", CONFIRMED, "A").keys()),
                         sorted(PROBED_FIELDNAMES))

    def test_the_address_from_the_popup_reaches_the_csv(self):
        row = probed("A", CONFIRMED, "A", city="Columbia", state="MD",
                     postcode="21046")
        self.assertEqual(row["City"], "Columbia")
        self.assertEqual(row["State"], "MD")
        self.assertEqual(row["ZIP"], "21046")

    def test_the_code_bottomline_showed_is_recorded_separately(self):
        row = probed("A", CODE_MISMATCH, "SOMEONE ELSE", code="V1",
                     popup_code="V999")
        self.assertEqual(row["ERP Vendor ID"], "V1")
        self.assertEqual(row["Bottomline Code"], "V999")


class ArchiveProbeCsv(unittest.TestCase):
    """The operator asked for each run's complete record to land in a NEW file
    (2026-09-09). Left alone, append_probed_rows() would put today's rows under
    last week's and the file would hold two different answers for one vendor
    with nothing to say which is current.

    It MOVES rather than deletes: a probe run is an hour at the screen and
    cannot be reproduced from anything on disk.
    """

    def setUp(self):
        self.folder = tempfile.mkdtemp()
        self.path = Path(self.folder) / "vendors_probed.csv"
        self.archive = Path(self.folder) / "archive"

    def tearDown(self):
        shutil.rmtree(self.folder, ignore_errors=True)

    def test_an_existing_file_is_moved_not_destroyed(self):
        self.path.write_text("Vendor Name\nACME\n", encoding="utf-8")
        moved = archive_probe_csv(self.path, self.archive, when="2026-09-09_1200")
        self.assertFalse(self.path.exists())
        self.assertTrue(moved.exists())
        self.assertIn("ACME", moved.read_text(encoding="utf-8"))

    def test_the_archived_name_carries_the_timestamp(self):
        self.path.write_text("x", encoding="utf-8")
        moved = archive_probe_csv(self.path, self.archive, when="2026-09-09_1200")
        self.assertEqual(moved.name, "vendors_probed_2026-09-09_1200.csv")

    def test_nothing_to_move_is_not_an_error(self):
        self.assertIsNone(archive_probe_csv(self.path, self.archive))
        self.assertFalse(self.archive.exists())

    def test_two_runs_in_the_same_minute_do_not_overwrite_each_other(self):
        """The counter is the whole point -- without it the second run of the
        minute would land on top of the first one's archive and destroy it."""
        self.path.write_text("first", encoding="utf-8")
        one = archive_probe_csv(self.path, self.archive, when="2026-09-09_1200")
        self.path.write_text("second", encoding="utf-8")
        two = archive_probe_csv(self.path, self.archive, when="2026-09-09_1200")
        self.assertNotEqual(one, two)
        self.assertEqual(one.read_text(encoding="utf-8"), "first")
        self.assertEqual(two.read_text(encoding="utf-8"), "second")

    def test_the_archive_folder_is_created_when_it_is_missing(self):
        self.path.write_text("x", encoding="utf-8")
        archive_probe_csv(self.path, self.archive)
        self.assertTrue(self.archive.is_dir())

    def test_a_fresh_run_then_writes_a_file_with_only_its_own_rows(self):
        """The behaviour the operator actually asked for, end to end."""
        append_probed_rows(self.path, [probed("OLD", CONFIRMED, "OLD")])
        archive_probe_csv(self.path, self.archive, when="2026-09-09_1200")
        append_probed_rows(self.path, [probed("NEW", CONFIRMED, "NEW")])
        with self.path.open(newline="", encoding="utf-8-sig") as handle:
            names = [r["Vendor Name"] for r in csv.DictReader(handle)]
        self.assertEqual(names, ["NEW"])


class WhatTheCodeBoxActuallyHeld(unittest.TestCase):
    """Added 2026-09-09 after the operator watched Bottomline complete a code as
    it was typed -- '762-' became '762-001' with the added part selected.

    Until now nothing ever read the box back, so a probe of the WRONG text
    looked identical in the log to a probe of the right one. Recording it costs
    one read and turns the question from a theory into something the next log
    settles by itself.
    """

    def test_an_attempt_carries_what_the_box_held(self):
        one = attempt("YARDI", "762", 1, {"name": "A", "code": "762-001"},
                      box="762-001")
        self.assertEqual(one["box"], "762-001")

    def test_the_box_defaults_to_empty_rather_than_missing(self):
        """Every older caller omits it, and none of them should have to guard
        for a key that is not there."""
        self.assertEqual(attempt("YARDI", "V1", 1)["box"], "")

    def test_it_is_trimmed_like_every_other_recorded_value(self):
        self.assertEqual(attempt("YARDI", "V1", 1, box="  V1  ")["box"], "V1")

    def test_it_reaches_the_csv_as_its_own_column(self):
        self.assertIn("Code Box Held", PROBED_FIELDNAMES)
        row = probed("A", CONFIRMED, "A", code="762", box="762-001")
        self.assertEqual(row["Code Box Held"], "762-001")
        self.assertEqual(row["ERP Vendor ID"], "762")

    def test_the_terminal_flags_a_box_that_holds_something_else(self):
        line = describe_probe(probed("ACME", NOT_FOUND, "", code="762",
                                     popup_code="", box="762-001"))
        self.assertIn("!box=", line)
        self.assertIn("762-001", line)

    def test_the_terminal_stays_quiet_when_the_box_is_correct(self):
        """A column that is always there stops being read."""
        line = describe_probe(probed("ACME", CONFIRMED, "ACME", code="V1",
                                     box="V1"))
        self.assertNotIn("!box=", line)

    def test_an_unread_box_is_not_reported_as_a_difference(self):
        """read_box_value() returns "" when the box cannot be read at all. That
        is 'we do not know', not 'the box was wrong', and must not be flagged."""
        line = describe_probe(probed("ACME", CONFIRMED, "ACME", code="V1",
                                     box=""))
        self.assertNotIn("!box=", line)

    def test_the_probe_reads_the_box_before_it_touches_the_popup(self):
        """Source-level: opening the popup moves focus, and anything that
        happens after that is no longer evidence of what the typing produced."""
        source = (HERE / "verify_master.py").read_text(encoding="utf-8")
        body = source[source.index("def probe_one_code("):]
        body = body[:body.index("\ndef ")]
        self.assertLess(body.index("read_box_value(page, box)"),
                        body.index("open_and_read_vendor_popup("))


class Checkpoint(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.mkdtemp()
        self.path = Path(self.folder) / ".probe_checkpoint.json"

    def tearDown(self):
        shutil.rmtree(self.folder, ignore_errors=True)

    def test_it_round_trips(self):
        write_checkpoint(self.path, 4, 39, 25)
        saved = read_checkpoint(self.path)
        self.assertEqual(saved["batches_done"], 4)
        self.assertEqual(saved["total_batches"], 39)
        self.assertEqual(saved["batch_size"], 25)

    def test_a_missing_checkpoint_reads_as_none(self):
        self.assertIsNone(read_checkpoint(self.path))

    def test_a_corrupt_checkpoint_reads_as_none_not_a_crash(self):
        """A checkpoint half-written by a crash is exactly when --resume matters
        most. A traceback here would send the operator back to row 1."""
        self.path.write_text("{ not json", encoding="utf-8")
        self.assertIsNone(read_checkpoint(self.path))

    def test_resume_starts_after_the_last_completed_batch(self):
        write_checkpoint(self.path, 4, 39, 25)
        self.assertEqual(resume_start_index(read_checkpoint(self.path)), 100)

    def test_resume_from_nothing_starts_at_the_beginning(self):
        self.assertEqual(resume_start_index(None), 0)

    def test_a_checkpoint_missing_its_fields_starts_at_the_beginning(self):
        self.assertEqual(resume_start_index({"batches_done": 4}), 0)


# ==========================================================================
# Comparing our master against Bottomline's answers
# ==========================================================================

class ReviewSheetFor(unittest.TestCase):
    """Every probed row lands on AT MOST ONE sheet. A row on two sheets gets
    settled twice and, worse, settled two different ways.
    """

    def test_a_clean_confirmation_needs_no_human(self):
        self.assertIsNone(review_sheet_for(probed("ACME LLC", CONFIRMED, "ACME LLC")))

    def test_a_different_spelling_goes_to_name_differs(self):
        self.assertEqual(
            review_sheet_for(probed("ACME LLC", CONFIRMED, "ACME LIMITED")),
            "Name Differs")

    def test_case_and_spacing_alone_are_not_a_difference(self):
        self.assertIsNone(review_sheet_for(probed("ACME  LLC", CONFIRMED, "acme llc")))

    def test_punctuation_alone_is_not_a_difference(self):
        self.assertIsNone(review_sheet_for(probed("ACME, LLC.", CONFIRMED, "ACME LLC")))

    def test_a_code_mismatch_gets_its_own_sheet(self):
        self.assertEqual(
            review_sheet_for(probed("A", CODE_MISMATCH, "SOMEONE ELSE",
                                    popup_code="V999")),
            "Code Mismatch")

    def test_wrong_package_goes_to_its_own_sheet(self):
        self.assertEqual(review_sheet_for(probed("A", WRONG_PACKAGE, "A")),
                         "Wrong Package")

    def test_not_found_goes_to_its_own_sheet(self):
        self.assertEqual(review_sheet_for(probed("A", NOT_FOUND)), "Not Found")

    def test_ambiguous_goes_to_its_own_sheet(self):
        self.assertEqual(review_sheet_for(probed("A", AMBIGUOUS, "A")), "Ambiguous")

    def test_a_learned_package_goes_to_its_own_sheet(self):
        self.assertEqual(
            review_sheet_for(probed("A", PACKAGE_LEARNED, "A", package="")),
            "Package Learned")

    def test_a_row_with_no_code_was_never_probed_so_needs_no_sheet(self):
        self.assertIsNone(review_sheet_for(probed("A", NO_CODE, code="")))

    def test_an_unverified_package_is_still_checked_for_its_name(self):
        self.assertIsNone(
            review_sheet_for(probed("ACME LLC", UNVERIFIED_PACKAGE, "ACME LLC")))
        self.assertEqual(
            review_sheet_for(probed("ACME LLC", UNVERIFIED_PACKAGE, "OTHER LLC")),
            "Name Differs")

    def test_a_code_mismatch_wins_over_a_name_difference(self):
        """One row, one sheet. The wrong code is the actionable half."""
        self.assertEqual(
            review_sheet_for(probed("ACME LLC", CODE_MISMATCH, "TOTALLY OTHER",
                                    popup_code="V999")),
            "Code Mismatch")

    def test_wrong_package_wins_over_a_name_difference(self):
        self.assertEqual(
            review_sheet_for(probed("ACME LLC", WRONG_PACKAGE, "TOTALLY OTHER")),
            "Wrong Package")


class NormaliseNameForCompare(unittest.TestCase):
    def test_case_is_ignored(self):
        self.assertEqual(normalise_name_for_compare("Acme"),
                         normalise_name_for_compare("ACME"))

    def test_runs_of_spaces_collapse(self):
        self.assertEqual(normalise_name_for_compare("A  B"),
                         normalise_name_for_compare("A B"))

    def test_punctuation_is_ignored(self):
        self.assertEqual(normalise_name_for_compare("ACME, LLC."),
                         normalise_name_for_compare("ACME LLC"))

    def test_none_and_blank_are_the_same(self):
        self.assertEqual(normalise_name_for_compare(None), "")
        self.assertEqual(normalise_name_for_compare(""), "")

    def test_a_real_difference_survives(self):
        self.assertNotEqual(normalise_name_for_compare("ACME LLC"),
                            normalise_name_for_compare("ACME LIMITED"))


class RouteToSheets(unittest.TestCase):
    def test_every_sheet_is_a_key_even_when_empty(self):
        routed = route_to_sheets([probed("A", CONFIRMED, "A")])
        self.assertEqual(sorted(routed.keys()), sorted(REVIEW_SHEETS))
        self.assertTrue(all(routed[s] == [] for s in REVIEW_SHEETS))

    def test_each_row_appears_at_most_once_across_all_sheets(self):
        rows = [probed("A", CONFIRMED, "DIFFERENT"), probed("B", NOT_FOUND),
                probed("C", AMBIGUOUS, "C"), probed("D", CONFIRMED, "D")]
        self.assertEqual(sum(len(v) for v in route_to_sheets(rows).values()), 3)

    def test_the_sheet_list_includes_code_mismatch(self):
        self.assertIn("Code Mismatch", REVIEW_SHEETS)


class WriteReviewWorkbook(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.mkdtemp()
        self.path = Path(self.folder) / "master_review.xlsx"

    def tearDown(self):
        shutil.rmtree(self.folder, ignore_errors=True)

    def _write(self, rows):
        write_review_workbook(self.path, route_to_sheets(rows),
                              summary_lines(rows, "test run"), "2026-09-09")

    def test_it_creates_every_sheet_plus_a_summary(self):
        from openpyxl import load_workbook
        self._write([probed("A", NOT_FOUND)])
        names = load_workbook(self.path).sheetnames
        for sheet in REVIEW_SHEETS:
            self.assertIn(sheet, names)
        self.assertIn("Summary", names)

    def test_a_settled_difference_disappears_on_the_next_run(self):
        """The sheets are REBUILT each run. If they were appended to instead,
        every difference a human fixed would sit in the report for ever."""
        from openpyxl import load_workbook
        self._write([probed("A", NOT_FOUND)])
        self.assertEqual(load_workbook(self.path)["Not Found"].max_row, 2)
        self._write([probed("A", CONFIRMED, "A")])
        self.assertEqual(load_workbook(self.path)["Not Found"].max_row, 1)

    def test_a_sheet_this_script_does_not_own_survives(self):
        from openpyxl import load_workbook
        self._write([probed("A", NOT_FOUND)])
        book = load_workbook(self.path)
        book.create_sheet("My Notes").append(["keep me"])
        book.save(self.path)
        self._write([probed("A", NOT_FOUND)])
        book = load_workbook(self.path)
        self.assertIn("My Notes", book.sheetnames)
        self.assertEqual(book["My Notes"]["A1"].value, "keep me")

    def test_codes_and_zips_are_text_so_excel_keeps_leading_zeros(self):
        """C/D/E are the three code columns, K is the ZIP. These are POSITIONS,
        so they moved when "Code Box Held" was inserted at E on 2026-09-09."""
        from openpyxl import load_workbook
        self._write([probed("A", NOT_FOUND, code="0012", postcode="04101")])
        sheet = load_workbook(self.path)["Not Found"]
        for column in ("C", "D", "E", "K"):
            self.assertEqual(sheet[f"{column}2"].number_format, "@")

    def test_the_headers_and_the_column_widths_stay_the_same_length(self):
        """The widths are applied by position. If a header is ever added without
        a width beside it, the last columns silently keep Excel's default and
        the report looks broken for no obvious reason."""
        from openpyxl import load_workbook
        self._write([probed("A", NOT_FOUND)])
        sheet = load_workbook(self.path)["Not Found"]
        widths = [sheet.column_dimensions[chr(ord("A") + n)].width
                  for n in range(len(REVIEW_HEADERS))]
        self.assertTrue(all(w for w in widths))

    def test_what_the_box_held_reaches_the_report(self):
        from openpyxl import load_workbook
        self._write([probed("A", CODE_MISMATCH, "SOMEONE ELSE", code="V1",
                            popup_code="V999", box="V1-001")])
        sheet = load_workbook(self.path)["Code Mismatch"]
        self.assertEqual(sheet["E1"].value, "Code Box Held")
        self.assertEqual(sheet["E2"].value, "V1-001")


class DescribeProbe(unittest.TestCase):
    """The operator asked to SEE what is captured rather than wait an hour and
    read a spreadsheet. It is also the fastest way to spot a run gone wrong: a
    column of ---- means the typing is outrunning the dropdown again, a column
    of MISM means the code box is not clearing.
    """

    def test_a_confirmation_shows_both_names_and_both_codes(self):
        line = describe_probe(probed("ACME LLC", CONFIRMED, "ACME LIMITED",
                                     code="V1"))
        self.assertIn("ACME LLC", line)
        self.assertIn("ACME LIMITED", line)
        self.assertIn("V1", line)
        self.assertIn("OK", line)

    def test_the_address_is_shown_when_there_is_one(self):
        line = describe_probe(probed("ACME", CONFIRMED, "ACME", city="Columbia",
                                     state="MD", postcode="21046"))
        self.assertIn("Columbia", line)
        self.assertIn("MD", line)
        self.assertIn("21046", line)

    def test_a_mismatch_is_marked_so_it_stands_out(self):
        line = describe_probe(probed("ACME", CODE_MISMATCH, "SOMEONE ELSE",
                                     code="V1", popup_code="V999"))
        self.assertIn("MISM", line)
        self.assertIn("V999", line)

    def test_a_not_found_says_so_without_inventing_a_name(self):
        line = describe_probe(probed("ACME", NOT_FOUND, "", code="V1",
                                     popup_code=""))
        self.assertIn("----", line)
        self.assertNotIn("->", line)

    def test_a_skipped_row_is_marked_skip(self):
        self.assertIn("skip", describe_probe(
            probed("ACME", NO_CODE, "", code="", popup_code="")))

    def test_every_verdict_has_a_mark_of_its_own(self):
        marks = [VERDICT_MARK[v] for v in
                 (CONFIRMED, CODE_MISMATCH, WRONG_PACKAGE, NOT_FOUND,
                  AMBIGUOUS, NO_CODE, PACKAGE_LEARNED, UNVERIFIED_PACKAGE)]
        self.assertEqual(len(marks), len(set(marks)))

    def test_a_very_long_vendor_name_does_not_run_away(self):
        line = describe_probe(probed("X" * 200, CONFIRMED, "Y" * 200))
        self.assertLess(len(line), 140)

    def test_it_never_raises_on_a_row_with_missing_keys(self):
        self.assertIsInstance(describe_probe({}), str)


class SummaryLines(unittest.TestCase):
    def test_it_counts_each_verdict(self):
        lines = summary_lines(
            [probed("A", CONFIRMED, "A"), probed("B", CONFIRMED, "B"),
             probed("C", NOT_FOUND)], "because")
        self.assertIn("CONFIRMED: 2", lines)
        self.assertIn("NOT_FOUND: 1", lines)
        self.assertIn("Rows probed: 3", lines)

    def test_the_phase_0_reason_is_always_carried_into_the_report(self):
        self.assertIn("Phase 0: because", summary_lines([], "because"))


# ==========================================================================
# The browser layer (fakes only; no browser ever opens)
# ==========================================================================

SELECTORS = {"fill": {"vendor_code": "#vendor-code", "vendor": "#vendor",
                      "integration_package": "#package"},
             "dropdown_item": ".x-boundlist-item:visible"}


class CountingLocator:
    def __init__(self, count):
        self._count = count

    def count(self):
        return self._count


class CountingPage:
    def __init__(self, count):
        self._count = count

    def locator(self, selector):
        return CountingLocator(self._count)


class RaisingPage:
    def locator(self, selector):
        raise RuntimeError("boom")


class IndexingFormIsOpen(unittest.TestCase):
    """The vendor CODE box exists only on the indexing screen, so it is the
    landmark that proves a form is really open. Without this check the run
    would happily probe a queue page or a half-loaded frame and read nothing
    for every code -- which looks exactly like "none of these codes exist".
    """

    def test_a_box_on_screen_means_the_form_is_open(self):
        self.assertTrue(indexing_form_is_open(CountingPage(1), SELECTORS))

    def test_no_box_means_no_form(self):
        self.assertFalse(indexing_form_is_open(CountingPage(0), SELECTORS))

    def test_an_unreadable_page_counts_as_no_form(self):
        self.assertFalse(indexing_form_is_open(RaisingPage(), SELECTORS))


class StickyBox:
    """A code box that behaves the way the real one did on 2026-09-09: the
    first clear leaves one character behind. `sticky_clears` says how many
    clears misbehave before it finally empties.
    """

    def __init__(self, value, sticky_clears=1):
        self.value = value
        self.sticky_clears = sticky_clears
        self.clears = 0
        self.selected = False

    def click(self):
        pass

    def press(self, key):
        if key == "Control+a":
            self.selected = True
            return
        if key in ("Delete", "Backspace"):
            if not self.selected:
                return
            self.clears += 1
            self.selected = False
            # The bug: an early clear leaves the first character behind.
            self.value = self.value[:1] if self.clears <= self.sticky_clears else ""

    def input_value(self):
        return self.value


class StickyLocator:
    def __init__(self, box):
        self.first = box


class StickyPage:
    def __init__(self, box):
        self.box = box
        self.waits = 0

    def locator(self, selector):
        return StickyLocator(self.box)

    def wait_for_timeout(self, ms):
        self.waits += 1


class ClearCodeBox(unittest.TestCase):
    """The bug that made the live run useless. 'ROL1234' was cleared, one
    character was left, and the NEXT code was typed onto it -- 'rROL5678' --
    which matches nothing. A perfectly good vendor is then reported as not
    found. The old code cleared once and assumed it had worked.
    """

    def test_an_already_empty_box_needs_one_pass(self):
        page = StickyPage(StickyBox("", sticky_clears=0))
        self.assertTrue(clear_code_box(page, "#code"))

    def test_a_clean_box_clears_first_time(self):
        box = StickyBox("ROL1234", sticky_clears=0)
        self.assertTrue(clear_code_box(StickyPage(box), "#code"))
        self.assertEqual(box.value, "")

    def test_a_leftover_character_is_noticed_and_cleared(self):
        box = StickyBox("ROL1234", sticky_clears=1)
        self.assertTrue(clear_code_box(StickyPage(box), "#code"))
        self.assertEqual(box.value, "")
        self.assertGreaterEqual(box.clears, 2)

    def test_a_very_stubborn_box_still_gets_cleared_within_the_attempts(self):
        box = StickyBox("ROL1234", sticky_clears=3)
        self.assertTrue(clear_code_box(StickyPage(box), "#code"))
        self.assertEqual(box.value, "")

    def test_a_box_that_never_clears_reports_failure(self):
        """It must NOT report success. probe_one_code() stops the run on False,
        because typing onto leftover text corrupts every probe after it too."""
        box = StickyBox("ROL1234", sticky_clears=99)
        self.assertFalse(clear_code_box(StickyPage(box), "#code"))

    def test_it_waits_between_attempts_so_extjs_can_notice(self):
        page = StickyPage(StickyBox("ROL1234", sticky_clears=1))
        clear_code_box(page, "#code")
        self.assertGreater(page.waits, 0)

    def test_a_raising_page_reports_failure_rather_than_pretending(self):
        self.assertFalse(clear_code_box(RaisingPage(), "#code"))


class ReadBoxValue(unittest.TestCase):
    def test_a_real_value_is_returned_trimmed(self):
        self.assertEqual(
            read_box_value(StickyPage(StickyBox("  ROL1234  ")), "#c"), "ROL1234")

    def test_an_empty_box_reads_as_empty(self):
        self.assertEqual(read_box_value(StickyPage(StickyBox("")), "#c"), "")

    def test_an_unreadable_box_reads_as_empty(self):
        self.assertEqual(read_box_value(RaisingPage(), "#c"), "")

    def test_a_missing_selector_reads_as_empty(self):
        self.assertEqual(read_box_value(StickyPage(StickyBox("X")), ""), "")


class ReadPauseChoice(unittest.TestCase):
    """Four actions at the pause. The default matters most: anything the
    operator types that is not understood must mean LOOK AGAIN, never skip.
    A typo that silently skipped a vendor would write a wrong answer and move
    on -- the exact failure the pause was added to stop.
    """

    def test_a_bare_enter_means_look_again(self):
        self.assertEqual(read_pause_choice(""), PAUSE_RETRY)

    def test_s_skips_this_vendor(self):
        self.assertEqual(read_pause_choice("s"), PAUSE_SKIP)

    def test_a_stops_asking_for_the_rest_of_the_run(self):
        self.assertEqual(read_pause_choice("a"), PAUSE_STOP_ASKING)

    def test_q_stops_the_run(self):
        self.assertEqual(read_pause_choice("q"), PAUSE_QUIT)

    def test_case_and_stray_spaces_are_forgiven(self):
        self.assertEqual(read_pause_choice("  S  "), PAUSE_SKIP)
        self.assertEqual(read_pause_choice("Q"), PAUSE_QUIT)

    def test_the_whole_word_works_too(self):
        self.assertEqual(read_pause_choice("skip"), PAUSE_SKIP)
        self.assertEqual(read_pause_choice("quit"), PAUSE_QUIT)

    def test_anything_unrecognised_looks_again_rather_than_skipping(self):
        self.assertEqual(read_pause_choice("x"), PAUSE_RETRY)
        self.assertEqual(read_pause_choice(None), PAUSE_RETRY)


class PauseMessage(unittest.TestCase):
    """What the operator reads at 11pm on row 400. It has to say which vendor,
    what was typed, and what came back -- without them scrolling or opening a
    file."""

    def test_it_names_the_vendor_and_the_code_that_was_typed(self):
        text = pause_message(vendor("ACME LLC", code="V1"),
                             miss("YARDI", "V1"))
        self.assertIn("ACME LLC", text)
        self.assertIn("V1", text)
        self.assertIn("YARDI", text)

    def test_a_popup_naming_someone_else_is_called_out_as_the_box_not_clearing(self):
        """The most valuable thing this message can do: turn the clearing bug
        into something the operator recognises instead of a mystery."""
        text = pause_message(vendor("ACME LLC", code="V1"),
                             other_vendor("YARDI", "V1", "V999", "PREVIOUS CO"))
        self.assertIn("PREVIOUS CO", text)
        self.assertIn("V999", text)
        self.assertIn("cleared", text)

    def test_it_says_when_nothing_could_be_read_at_all(self):
        text = pause_message(vendor("ACME LLC", code="V1"), miss("YARDI", "V1"))
        self.assertIn("showed nothing", text)

    def test_it_reassures_that_nothing_was_committed(self):
        text = pause_message(vendor("A", code="V1"), miss())
        self.assertIn("Nothing has been committed", text)

    def test_a_row_with_missing_keys_does_not_crash_it(self):
        self.assertIsInstance(pause_message({}, miss()), str)


class FakeElement:
    def __init__(self):
        self.clicks = 0

    def click(self):
        self.clicks += 1


class FakeHandle:
    def __init__(self, element):
        self._element = element

    def as_element(self):
        return self._element


class ScriptedPopupPage:
    """A page whose Vendor Address popup shows something different each time it
    is opened -- which is exactly what the screen looks like when the operator
    fixes something between two clicks.

    `texts` is what the popup shows on the 1st, 2nd, ... opening. The last one
    repeats. Escape closes it, the same way close_vendor_popup() expects.
    """

    def __init__(self, texts, matches=1):
        self.texts = list(texts) or [""]
        self.matches = matches
        self.opened = 0
        self.showing = ""
        self.element = FakeElement()
        self.keyboard = self

    def locator(self, selector):
        return CountingLocator(self.matches)

    def evaluate_handle(self, script, argument):
        # Finding the Vendor Address control; clicking it opens the popup.
        self.showing = self.texts[min(self.opened, len(self.texts) - 1)]
        self.opened += 1
        return FakeHandle(self.element)

    def evaluate(self, script, argument):
        return self.showing

    def wait_for_timeout(self, ms):
        pass

    def press(self, key):
        if key == "Escape":
            self.showing = ""


class Answers:
    """Stands in for input() at the pause prompt."""

    def __init__(self, *replies):
        self.replies = list(replies)
        self.asked = 0

    def __call__(self, prompt=""):
        self.asked += 1
        return self.replies.pop(0) if self.replies else "s"


class PauseUntilLoaded(unittest.TestCase):
    """The operator's instruction, 2026-09-09: "if no vendor loads, tell the
    user to fix it and then it will resume in terminal and script will continue
    from there -- clicking vendor address and fetching the needed data".

    Everything here turns on ONE rule: the retry re-reads the popup and never
    retypes the code. Retyping would undo the fix the run just stopped to let
    the operator make.
    """

    def setUp(self):
        self.folder = tempfile.mkdtemp()
        self.log = Path(self.folder) / "probe.txt"
        self.pause_state = {"asking": True}
        self.row = vendor("ACME LLC", code="V1", package="YARDI (2)")

    def tearDown(self):
        shutil.rmtree(self.folder, ignore_errors=True)

    def _run(self, texts, answers, matches=1, failed=None):
        page = ScriptedPopupPage(texts, matches)
        ask = Answers(*answers)
        result = pause_until_loaded(page, SELECTORS,
                                    failed or miss("YARDI", "V1"),
                                    self.log, self.pause_state, self.row, ask)
        return page, ask, result

    def test_enter_looks_again_and_the_fixed_vendor_is_returned(self):
        page, ask, result = self._run(
            ["Vendor Address\nACME LLC (V1)\nColumbia, MD 21046"], [""])
        self.assertIsNotNone(result)
        self.assertEqual(result["popup_code"], "V1")
        self.assertEqual(result["popup_name"], "ACME LLC")
        self.assertEqual(ask.asked, 1)

    def test_the_address_from_the_fixed_read_is_captured_too(self):
        _, _, result = self._run(
            ["Vendor Address\nACME LLC (V1)\nColumbia, MD 21046"], [""])
        self.assertEqual(result["city"], "Columbia")
        self.assertEqual(result["state"], "MD")
        self.assertEqual(result["zip"], "21046")

    def test_it_keeps_asking_until_the_vendor_finally_loads(self):
        page, ask, result = self._run(
            ["Vendor Address\nSOMEONE ELSE (V999)",
             "",
             "Vendor Address\nACME LLC (V1)"],
            ["", "", ""])
        self.assertEqual(ask.asked, 3)
        self.assertEqual(result["popup_code"], "V1")

    def test_skip_returns_nothing_so_the_honest_failure_stands(self):
        _, ask, result = self._run(["Vendor Address\nACME LLC (V1)"], ["s"])
        self.assertIsNone(result)
        self.assertEqual(ask.asked, 1)

    def test_stop_asking_turns_the_pause_off_for_the_whole_run(self):
        _, _, result = self._run(["Vendor Address\nACME LLC (V1)"], ["a"])
        self.assertIsNone(result)
        self.assertFalse(self.pause_state["asking"])

    def test_once_pausing_is_off_it_never_asks_again(self):
        self.pause_state["asking"] = False
        _, ask, result = self._run(["Vendor Address\nACME LLC (V1)"], [])
        self.assertIsNone(result)
        self.assertEqual(ask.asked, 0)

    def test_quit_stops_the_run_so_finished_batches_are_kept(self):
        """RuntimeError is what main() already catches to save every completed
        batch and print the --resume line."""
        with self.assertRaises(RuntimeError):
            self._run(["Vendor Address\nACME LLC (V1)"], ["q"])

    def test_the_quit_message_names_the_vendor_it_stopped_on(self):
        try:
            self._run(["Vendor Address\nACME LLC (V1)"], ["q"])
        except RuntimeError as error:
            self.assertIn("ACME LLC", str(error))
            self.assertIn("V1", str(error))

    def test_a_vendor_that_did_load_is_never_paused_on(self):
        page, ask, result = self._run(
            [""], [], failed=hit("YARDI", "V1", "ACME LLC"))
        self.assertEqual(ask.asked, 0)
        self.assertEqual(result["popup_name"], "ACME LLC")

    def test_the_retry_clicks_the_address_control_and_does_not_retype(self):
        """The whole point. probe_one_code() types; the retry must only click
        and read, or it would wipe out whatever the operator just fixed."""
        page, _, _ = self._run(["Vendor Address\nACME LLC (V1)"], [""])
        self.assertEqual(page.element.clicks, 1)
        self.assertEqual(page.opened, 1)

    def test_the_pause_is_written_into_the_log(self):
        self._run(["Vendor Address\nACME LLC (V1)"], [""])
        self.assertIn("PAUSED", self.log.read_text(encoding="utf-8"))


class PauseIsWiredInOncePerRow(unittest.TestCase):
    """Source-level, because the WHERE matters as much as the what. The pause
    must fire once per master row, AFTER every package that row is tried under
    -- never after a single probe. A miss under the claimed package is not a
    failure while the other package is still to come, and pausing there would
    ask the operator to fix something the script was about to answer itself.
    """

    def setUp(self):
        """The CODE of probe_batch, with its docstring cut off.

        The docstring names pause_until_loaded() while explaining itself, and
        that mention sits above every line of real code -- so a search over the
        whole function finds the prose, not the call, and the ordering tests
        below would pass no matter where the call actually went.
        """
        source = (HERE / "verify_master.py").read_text(encoding="utf-8")
        body = source[source.index("def probe_batch("):]
        body = body[:body.index("\n# Ctrl+C")]
        self.body = body[body.index('"""', body.index('"""') + 3) + 3:]

    def test_the_pause_happens_after_the_other_package_has_been_tried(self):
        self.assertLess(self.body.rindex("probe_one_code("),
                        self.body.index("pause_until_loaded("))

    def test_it_only_pauses_when_no_attempt_hit(self):
        self.assertIn("not any(_is_hit(one) for one in attempts)", self.body)

    def test_the_verdict_is_decided_after_the_pause_not_before(self):
        """rindex, not index: classify_probe() is also called near the top of
        the loop for a row with no code at all, and that one legitimately comes
        before the pause. It is the LAST call -- the real verdict -- that has to
        see whatever the operator fixed."""
        self.assertLess(self.body.index("pause_until_loaded("),
                        self.body.rindex("classify_probe("))

    def test_phase_0_never_pauses(self):
        """A calibration miss is a RESULT -- half the table Phase 0 is built to
        produce. Stopping to 'fix' it would destroy the measurement."""
        source = (HERE / "verify_master.py").read_text(encoding="utf-8")
        phase0 = source[source.index("def run_phase0("):]
        phase0 = phase0[:phase0.index("\ndef ")]
        self.assertNotIn("pause_until_loaded(", phase0)


class PackageAlreadySelected(unittest.TestCase):
    """Every row goes package -> code -> read, in that order, because Bottomline
    scopes what a box resolves to the package currently selected. This is the
    check that decides whether the package box needs touching for this row.
    """

    def test_the_same_package_needs_no_reselecting(self):
        self.assertTrue(package_already_selected("YARDI", "YARDI"))

    def test_the_bracketed_form_from_the_master_matches_the_plain_one(self):
        """'YARDI (2)' in vendors.csv and 'YARDI' on screen are one package.
        Without this every row would re-select needlessly -- or worse, decide
        the box had drifted when it had not."""
        self.assertTrue(package_already_selected("YARDI", "YARDI (2)"))
        self.assertTrue(package_already_selected("INTACCT (1)", "INTACCT"))

    def test_case_and_spacing_are_forgiven(self):
        self.assertTrue(package_already_selected(" yardi ", "YARDI"))

    def test_a_different_package_must_be_selected(self):
        self.assertFalse(package_already_selected("INTACCT", "YARDI"))

    def test_an_empty_box_is_never_treated_as_already_selected(self):
        """An unreadable or empty package box must mean "select it", never
        "assume it is fine" -- assuming would probe under whatever the site
        happens to have."""
        self.assertFalse(package_already_selected("", "YARDI"))
        self.assertFalse(package_already_selected(None, "YARDI"))


class PackageIsCheckedOnScreenEveryRow(unittest.TestCase):
    """2026-09-09: ensure_package() used to skip its work whenever its own
    `state` variable said the right package was chosen. That is only safe while
    nothing else touches the form -- and opening and closing the Vendor Address
    popup does. A stale belief there probes every following row under the WRONG
    package, and the answers still look reasonable.
    """

    def setUp(self):
        self.source = (HERE / "verify_master.py").read_text(encoding="utf-8")

    def test_it_reads_the_package_box_rather_than_trusting_state(self):
        body = self.source[self.source.index("def ensure_package("):]
        body = body[:body.index("\ndef ")]
        self.assertIn("_read_selected_package(page, selectors)", body)

    def test_it_no_longer_returns_early_on_the_cached_value_alone(self):
        body = self.source[self.source.index("def ensure_package("):]
        body = body[:body.index("\ndef ")]
        self.assertNotIn('if state.get("package") == package:\n        return True',
                         body)

    def test_it_proves_the_pick_landed_before_reporting_success(self):
        """Clicking an ExtJS dropdown item can fail silently -- the 2026-08-14
        invisible-item failure. A package that did not take would send every
        code after it to the wrong accounting system's list."""
        body = self.source[self.source.index("def ensure_package("):]
        body = body[:body.index("\ndef ")]
        after_pick = body[body.index("pick_only_match(page, items)"):]
        self.assertIn("_read_selected_package", after_pick)

    def test_the_package_is_selected_before_the_code_is_typed(self):
        """Order matters and is easy to break by moving one line."""
        body = self.source[self.source.index("def probe_batch("):]
        body = body[:body.index("\n# Ctrl+C")]
        self.assertLess(body.index("ensure_package("), body.index("probe_one_code("))


class SelectorKeysMatchLoadSelectors(unittest.TestCase):
    """The bug that killed the first live run, 2026-09-09.

    load_selectors() does not hand back selectors.json as-is: it RENAMES
    "dropdown_item_selector" to "dropdown_item". verify_master.py used the JSON
    spelling, so it raised KeyError on the first package switch -- after the
    login, after Phase 0 had already announced itself.
    """

    def setUp(self):
        from invoice_recorder import load_selectors
        self.selectors = load_selectors()
        self.source = (HERE / "verify_master.py").read_text(encoding="utf-8")

    def test_the_dropdown_key_is_the_one_load_selectors_actually_returns(self):
        self.assertIn("dropdown_item", self.selectors)
        self.assertNotIn("dropdown_item_selector", self.selectors)

    def test_the_prober_never_looks_up_the_json_spelling(self):
        self.assertNotIn('selectors["dropdown_item_selector"]', self.source)

    def test_every_box_the_prober_types_into_really_exists(self):
        for key in ("vendor_code", "vendor", "integration_package"):
            self.assertIn(key, self.selectors["fill"])

    def test_the_queue_navigator_is_no_longer_used(self):
        """Changed 2026-09-09: the operator opens an invoice. The prober does
        not care which one, so reading up to 50 queue pages to choose was the
        slowest part of a run and bought nothing."""
        self.assertNotIn("from queue_navigator import", self.source)
        self.assertNotIn("read_queue(", self.source)


class NothingCanWriteToBottomline(unittest.TestCase):
    """The reason this script is safe to point at a live invoice. It opens no
    invoice itself, types into boxes, reads a read-only popup, and leaves. If
    any of these names ever appears in the source that guarantee is gone -- so
    the SOURCE ITSELF is what is under test here.
    """

    def setUp(self):
        self.source = (HERE / "verify_master.py").read_text(encoding="utf-8")

    def test_it_never_names_a_button_that_commits(self):
        for forbidden in ("Convert to Invoice", "Save as Exception",
                          '"Save"', "'Save'", "never_click"):
            self.assertNotIn(forbidden, self.source,
                             f"verify_master.py must never reference {forbidden}")

    def test_delete_only_ever_appears_as_a_KEYSTROKE(self):
        """"Delete" is both a keyboard key and one of Bottomline's committing
        buttons. Clearing a box needs the key; nothing may ever click the
        button. So every mention must be a .press(), never a click."""
        for line in self.source.splitlines():
            if '"Delete"' in line or "'Delete'" in line:
                self.assertIn(".press(", line, f"Delete used outside a keystroke: {line}")

    def test_it_never_marks_a_spreadsheet_row_done(self):
        self.assertNotIn("mark_row_done", self.source)

    def test_it_never_writes_the_vendor_master(self):
        """P5: the master does not change without a human. A prober that
        'corrected' vendors.csv would destroy the only record of what we
        believed before the probe ran."""
        self.assertNotIn("write_master", self.source)
        for writing in ("open(VENDORS_CSV", "VENDORS_CSV.open",
                        "VENDORS_CSV.write"):
            self.assertNotIn(writing, self.source)

    def test_the_vendor_dropdown_is_never_picked_from(self):
        """P3 still holds. The user proved 2026-09-09 that the Vendor Address
        popup shows the right vendor from a TYPED code with nothing selected --
        which is exactly why no dropdown click was ever needed."""
        self.assertEqual(self.source.count("pick_only_match("), 1)
        for line in self.source.splitlines():
            if "pick_only_match(" in line:
                self.assertIn("items", line)

    def test_every_click_in_the_script_is_accounted_for(self):
        """Exactly four things are clicked and each is named here:

          1. the Vendor Address button      -- opens a read-only popup
          2. that popup's close button      -- closes it again
          3. `box.click()`                  -- focusing a text box before
                                               typing into it, which commits
                                               nothing
          4. `tool.click()` inside           -- ExtJS's own close tool on the
             _CLOSE_POPUP_JS                   Vendor Address window, clicked
                                               from JavaScript

        plus pick_only_match(), which clicks the Integration Package dropdown
        item and is counted separately below. If a fifth click ever appears,
        this fails and somebody has to justify it.

        WHY (4) EARNS ITS PLACE, added 2026-09-09 after the run died: while the
        popup is open ExtJS lays a modal mask over the page, and a real mouse
        click is hit-tested against that mask and swallowed -- which is what
        produced a 30-second Playwright timeout and a raw traceback. A JS
        .click() dispatches straight at the element and is not hit-tested, so it
        goes through. It is scoped to `[auto_id="vendor_address_window"]` and to
        an element matching `.x-tool-close`, so it can only ever close that one
        popup -- it cannot reach a committing button anywhere on the form.
        """
        allowed = ("POPUP_BUTTON", "POPUP_CLOSE", "box.click()",
                   "element.click()", "tool.click()")
        clicking = [line.strip() for line in self.source.splitlines()
                    if ".click()" in line and not line.strip().startswith("#")]
        for line in clicking:
            self.assertTrue(any(name in line for name in allowed),
                            f"unexplained click: {line}")

    def test_the_address_button_is_not_found_by_role_alone(self):
        """The bug found 2026-09-09: get_by_role("button", ...) never matched,
        so the click silently did nothing and no popup opened. content.js:49
        says why -- on the real site it "isn't a strict <button> tag". The text
        match must be tried FIRST, with role only as a fallback."""
        self.assertIn("_FIND_BUTTON_JS", self.source)
        text_match = self.source.index("text-match")
        role_button = self.source.index('get_by_role("button", name=POPUP_BUTTON')
        self.assertLess(text_match, role_button,
                        "the text match must be attempted before role=button")

    def test_it_does_not_borrow_the_recorders_resume_line(self):
        """log_interrupted() prints 'python invoice_recorder.py --start-row N'.
        The prober resumes by BATCH from its checkpoint, so borrowing that
        function would print a command that does the wrong thing."""
        self.assertNotIn("    log_interrupted,", self.source)
        self.assertIn("log_probe_interrupted()", self.source)
        self.assertIn("--resume", self.source)


if __name__ == "__main__":
    unittest.main(verbosity=2)
