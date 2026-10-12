"""
Build the Excel test log from test_log.csv
==========================================
The CSV is the record; this rebuilds a readable Excel copy of it. Run it after
adding a row. Reads nothing else, writes nothing else, touches no invoice.

    python build_test_log.py
"""

import csv
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = Path(__file__).resolve().parent
CSV_PATH = HERE / "test_log.csv"
XLSX_PATH = HERE / "ap_splitter_test_log.xlsx"

# Verdict colours: green pass, red fail, amber partial.
FILLS = {
    "PASS": PatternFill("solid", fgColor="C6EFCE"),
    "FAIL": PatternFill("solid", fgColor="FFC7CE"),
    "PARTIAL": PatternFill("solid", fgColor="FFEB9C"),
}
FONTS = {
    "PASS": Font(color="006100", bold=True),
    "FAIL": Font(color="9C0006", bold=True),
    "PARTIAL": Font(color="9C6500", bold=True),
}
WIDTHS = [6, 12, 11, 34, 7, 24, 24, 10, 55, 55, 75, 12, 10]


def main():
    if not CSV_PATH.exists():
        sys.exit(f"Not found: {CSV_PATH}")

    with open(CSV_PATH, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    if not rows:
        sys.exit("test_log.csv is empty")

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Log"
    for row in rows:
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="4F81BD")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    verdict_col = rows[0].index("verdict") + 1
    wrap = Alignment(wrap_text=True, vertical="top")
    for r in range(2, ws.max_row + 1):
        verdict = str(ws.cell(row=r, column=verdict_col).value or "").upper()
        if verdict in FILLS:
            cell = ws.cell(row=r, column=verdict_col)
            cell.fill, cell.font = FILLS[verdict], FONTS[verdict]
            cell.alignment = Alignment(horizontal="center", vertical="top")
        for c in range(1, ws.max_column + 1):
            if c != verdict_col:
                ws.cell(row=r, column=c).alignment = wrap

    for width, col in zip(WIDTHS, "ABCDEFGHIJKLM"):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # A short summary block under the table.
    verdicts = [str(ws.cell(row=r, column=verdict_col).value or "").upper()
                for r in range(2, ws.max_row + 1)]
    start = ws.max_row + 2
    ws.cell(row=start, column=1, value="Runs").font = Font(bold=True)
    ws.cell(row=start, column=2, value=len(verdicts))
    for i, name in enumerate(["PASS", "FAIL", "PARTIAL"], start=1):
        ws.cell(row=start + i, column=1, value=name).font = Font(bold=True)
        ws.cell(row=start + i, column=2, value=verdicts.count(name))

    try:
        wb.save(XLSX_PATH)
    except PermissionError:
        sys.exit(f"\n{XLSX_PATH.name} is open in Excel. Close it and run again.\n"
                 f"Nothing is lost - test_log.csv still holds every row.\n")

    print(f"Wrote {XLSX_PATH}")
    print(f"  {len(verdicts)} run(s):  "
          f"{verdicts.count('PASS')} pass, {verdicts.count('FAIL')} fail, "
          f"{verdicts.count('PARTIAL')} partial")


if __name__ == "__main__":
    main()
