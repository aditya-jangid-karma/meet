"""
Invoice Boundary Detection Engine
=================================
Reads multi-invoice PDFs and PREDICTS which page each new invoice starts on,
then writes a reviewer-friendly Excel report. It NEVER splits the PDFs.

Run modes:
    python invoice_boundary_detector.py            → read PDFs, predict, write Excel
    python invoice_boundary_detector.py --score    → read the reviewed Excel, print accuracy
    python invoice_boundary_detector.py --selftest  → run built-in logic checks

Requirements:
    pip install PyMuPDF pytesseract Pillow openpyxl tqdm
    Tesseract-OCR installed (path set in config.json)
"""

import os
import io
import re
import sys
import json
import logging
import argparse
from pathlib import Path
from datetime import datetime

# Third-party (installed via pip). Imported lazily inside functions that need
# them so that --selftest can run pure-logic checks without them installed.
try:
    from tqdm import tqdm
except ImportError:
    def tqdm(iterable, **kwargs):  # no-op fallback if tqdm isn't installed
        return iterable

SCRIPT_DIR = Path(__file__).resolve().parent


# ===========================================================================
# SECTION 1 — CONFIG
# Plain English: this reads config.json, the ONE settings file. It holds the
# folder paths, where Tesseract lives, the OCR image quality (dpi), the score
# needed to call a page a "new invoice", the points each rule is worth, and
# the list of known vendor names. Non-engineers tune behavior HERE, not in code.
# ===========================================================================

def load_config(path: str = "config.json") -> dict:
    config_path = SCRIPT_DIR / path
    with open(config_path, "r", encoding="utf-8") as f:
        return json.load(f)


# ===========================================================================
# SECTION 2 — READER (OCR)
# Plain English: opens each PDF, turns every page into a 300-dpi image
# (PyMuPDF), and reads the text off that image with Tesseract OCR. We OCR
# EVERY page (even digital ones) for uniform behavior. If a page cannot be
# read, we DO NOT silently skip it — we record WHY (the error message) so the
# reviewer sees the reason later.
# ===========================================================================

def collect_pdfs(folder: Path) -> list:
    """Return a sorted list of all *.pdf paths directly under `folder`."""
    if not folder.exists():
        return []
    return sorted(folder.glob("*.pdf"))


def read_pdf_pages(pdf_path: Path, config: dict) -> list:
    """OCR every page of one PDF. Returns a list of PageRecord dicts."""
    import fitz  # PyMuPDF
    import pytesseract
    from PIL import Image

    pytesseract.pytesseract.tesseract_cmd = config["tesseract_cmd"]
    dpi = config["ocr"]["dpi"]
    pdf_name = pdf_path.name
    records = []

    try:
        doc = fitz.open(str(pdf_path))
    except Exception as e:
        # Whole file unreadable: one record explaining why.
        return [{
            "pdf_name": pdf_name, "page_number": 1, "text": "",
            "read_status": "failed", "read_reason": f"could not open PDF: {e}",
        }]

    for page_index, page in enumerate(doc, start=1):
        try:
            pix = page.get_pixmap(dpi=dpi)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            text = pytesseract.image_to_string(img)
            records.append({
                "pdf_name": pdf_name, "page_number": page_index, "text": text,
                "read_status": "ok", "read_reason": None,
            })
        except Exception as e:
            records.append({
                "pdf_name": pdf_name, "page_number": page_index, "text": "",
                "read_status": "failed",
                "read_reason": f"OCR failed on page {page_index}: {e}",
            })
    doc.close()
    return records


# ===========================================================================
# SECTION 3 — FEATURES
# Plain English: given the text of one page, pull out the useful fields using
# pattern matching (regular expressions). These patterns are adapted from the
# earlier ocr_extract.py. Anything not found comes back as None.
# ===========================================================================

_DOC_NUMBER_PATTERNS = [
    re.compile(r"(?:invoice|inv)\s*(?:#|no\.?|number)\s*[:\-]?\s*([A-Z0-9][\w\-/]{1,30})", re.I),
    re.compile(r"(?:statement)\s*(?:#|no\.?|number)\s*[:\-]?\s*([A-Z0-9][\w\-/]{1,30})", re.I),
    re.compile(r"(?:bill|receipt|credit\s*memo|debit\s*memo)\s*(?:#|no\.?|number)\s*[:\-]?\s*([A-Z0-9][\w\-/]{1,30})", re.I),
]
_PO_PATTERN = re.compile(r"(?:purchase\s*order|po)\s*(?:#|no\.?|number)\s*[:\-]?\s*([A-Z0-9][\w\-/]{1,30})", re.I)
_DATE_PATTERNS = [
    re.compile(r"(?:invoice\s*date|statement\s*date|bill\s*date|dated|date)\s*[:\-]?\s*(\d{1,2}[/\-\.]\d{1,2}[/\-\.]\d{2,4})", re.I),
    re.compile(r"(?:invoice\s*date|statement\s*date|bill\s*date|dated|date)\s*[:\-]?\s*([A-Za-z]+\.?\s+\d{1,2},?\s+\d{2,4})", re.I),
]
_DUE_DATE_PATTERNS = [
    re.compile(r"(?:due\s*date|payment\s*due|pay\s*by|due\s*on)\s*[:\-]?\s*(\d{1,2}[/\-\.]\d{1,2}[/\-\.]\d{2,4})", re.I),
]
_AMOUNT_PATTERNS = [
    re.compile(r"(?:total\s*(?:due|amount|balance)?|amount\s*due|balance\s*due|grand\s*total|net\s*amount)\s*[:\-]?\s*\$?\s*([\d,]+\.\d{2})", re.I),
]
_BILL_TO_PATTERNS = [
    re.compile(r"(?:bill\s*to|sold\s*to|ship\s*to|customer|client|attn)\s*[:\-]?\s*(.+?)(?:\n|$)", re.I),
]
_CURRENCY_PATTERN = re.compile(r"\b(USD|CAD|EUR|GBP)\b|(\$)")


def _first(patterns: list, text: str) -> str | None:
    for p in patterns:
        m = p.search(text)
        if m:
            return m.group(1).strip()
    return None


def find_doc_number(text: str) -> str | None:
    return _first(_DOC_NUMBER_PATTERNS, text)


def extract_features(text: str) -> dict:
    currency = None
    m = _CURRENCY_PATTERN.search(text)
    if m:
        currency = m.group(1) if m.group(1) else "USD"
    return {
        "invoice_number": find_doc_number(text),
        "date": _first(_DATE_PATTERNS, text),
        "due_date": _first(_DUE_DATE_PATTERNS, text),
        "vendor_raw": None,  # filled by VENDOR section (Task 4)
        "bill_to": _first(_BILL_TO_PATTERNS, text),
        "total": _first(_AMOUNT_PATTERNS, text),
        "po_number": (_PO_PATTERN.search(text).group(1).strip() if _PO_PATTERN.search(text) else None),
        "currency": currency,
    }


# ===========================================================================
# SECTION 4 — VENDOR
# Plain English: figure out WHO the invoice is from. Vendor drives file naming
# and a key rule, so we try several methods, most reliable first:
#   1) known-vendor list (config)   2) top-of-page company line
#   3) email/website domain          4) old "From:/Vendor:" label
# We attach a confidence and never trust a guess silently.
# ===========================================================================

_LABEL_VENDOR = re.compile(r"(?:from|vendor|supplier|remit\s*to|sold\s*by|billed?\s*by)\s*[:\-]?\s*(.+?)(?:\n|$)", re.I)
_DOMAIN = re.compile(r"[\w.\-]+@([A-Za-z0-9\-]+)\.(?:com|net|org|co)", re.I)
_COMPANY_SUFFIX = re.compile(r"\b([A-Z][A-Za-z&'.\- ]{2,40}?(?:LLC|Inc|Incorporated|Co|Company|Corp|Ltd|Supply|Services|Plumbing|Electric))\b")
_ILLEGAL_FILENAME = re.compile(r'[\\/:*?"<>|]')


def sanitize_filename(name: str) -> str:
    return _ILLEGAL_FILENAME.sub("_", name).strip()


def detect_vendor(text: str, config: dict) -> dict:
    known = config.get("known_vendors", [])

    # 1) Known-vendor list (case-insensitive substring).
    lowered = text.lower()
    for name in known:
        if name.lower() in lowered:
            return {"vendor_name": name, "confidence": "high", "source": "known-list"}

    # 2) Top-of-page company line (first ~8 non-empty lines, look for a suffix).
    top = "\n".join([ln for ln in text.splitlines() if ln.strip()][:8])
    m = _COMPANY_SUFFIX.search(top)
    if m:
        return {"vendor_name": m.group(1).strip(), "confidence": "med", "source": "top-of-page"}

    # 3) Email/website domain.
    m = _DOMAIN.search(text)
    if m:
        return {"vendor_name": m.group(1).strip(), "confidence": "med", "source": "domain"}

    # 4) Label fallback.
    m = _LABEL_VENDOR.search(text)
    if m and len(m.group(1).strip()) >= 3:
        return {"vendor_name": m.group(1).strip(), "confidence": "low", "source": "label"}

    return {"vendor_name": None, "confidence": "low", "source": "none"}


# ===========================================================================
# SECTION 5 — RULES
# Plain English: each rule is a small, independent check comparing the current
# page to the previous one. If it "fires", it returns points (+/-) and a plain
# reason. Points come from config so non-engineers can tune them. Positive
# points push toward "a new invoice starts here" (Break); negative points push
# toward "same invoice" (Continue).
# ===========================================================================

_HEADER_HINT = re.compile(r"\b(invoice|statement|bill\s*to|remit\s*to|invoice\s*date)\b", re.I)
_CONTINUATION_HINT = re.compile(r"\b(continued|page\s*\d+\s*of\s*\d+|carried\s*forward)\b", re.I)


def rule_invoice_number_changed(prev, curr, w):
    a = prev["features"].get("invoice_number")
    b = curr["features"].get("invoice_number")
    if a and b and a != b:
        return w["invoice_number_changed"], f"Invoice number changed ({a} -> {b})"
    return None

def rule_vendor_changed(prev, curr, w):
    a = (prev.get("vendor") or {}).get("vendor_name")
    b = (curr.get("vendor") or {}).get("vendor_name")
    if a and b and a != b:
        return w["vendor_changed"], f"Vendor changed ({a} -> {b})"
    return None

def rule_invoice_header_detected(prev, curr, w):
    if _HEADER_HINT.search(curr.get("text", "")):
        return w["invoice_header_detected"], "Invoice header keywords on this page"
    return None

def rule_same_vendor(prev, curr, w):
    a = (prev.get("vendor") or {}).get("vendor_name")
    b = (curr.get("vendor") or {}).get("vendor_name")
    if a and b and a == b:
        return w["same_vendor"], f"Same vendor as previous page ({a})"
    return None

def rule_continuation_pattern(prev, curr, w):
    txt = curr.get("text", "")
    if _CONTINUATION_HINT.search(txt) or (not curr["features"].get("invoice_number") and not _HEADER_HINT.search(txt)):
        return w["continuation_pattern"], "Looks like a continuation page"
    return None

RULES = [
    rule_invoice_number_changed,
    rule_vendor_changed,
    rule_invoice_header_detected,
    rule_same_vendor,
    rule_continuation_pattern,
]


# ===========================================================================
# SECTION 6 — CONFIDENCE
# Plain English: add up all the rule points for a page. If the total reaches
# the threshold in config, we say a NEW invoice starts here (Break); otherwise
# it CONTINUES the current one. Confidence = how far the total is from the line.
# ===========================================================================

def _confidence_from_margin(margin: int) -> str:
    if margin >= 40:
        return "high"
    if margin >= 15:
        return "med"
    return "low"


def score_boundary(prev: dict, curr: dict, config: dict) -> dict:
    w = config["rule_weights"]
    threshold = config["thresholds"]["break_at_score"]
    total, reasons = 0, []
    for rule in RULES:
        out = rule(prev, curr, w)
        if out is not None:
            pts, why = out
            total += pts
            reasons.append(f"{why} ({pts:+d})")
    decision = "Break" if total >= threshold else "Continue"
    return {
        "page_number": curr["page_number"],
        "decision": decision,
        "confidence": _confidence_from_margin(abs(total - threshold)),
        "total_score": total,
        "fired_rules": reasons,
    }


def predict_boundaries(pages: list, config: dict) -> list:
    """Page 1 always starts an invoice (Break). Pages 2..n are scored."""
    preds = []
    for i, page in enumerate(pages):
        if i == 0:
            preds.append({
                "page_number": page["page_number"], "decision": "Break",
                "confidence": "high", "total_score": config["thresholds"]["break_at_score"],
                "fired_rules": ["First page of the PDF"],
            })
        else:
            preds.append(score_boundary(pages[i - 1], page, config))
    return preds


# ---- Pipeline glue: turn raw PageRecords into enriched pages, then summarize.
def build_pages(records: list, config: dict) -> list:
    pages = []
    for r in records:
        feats = extract_features(r["text"])
        vend = detect_vendor(r["text"], config)
        feats["vendor_raw"] = vend["vendor_name"]
        pages.append({
            "page_number": r["page_number"], "text": r["text"],
            "read_status": r["read_status"], "read_reason": r["read_reason"],
            "features": feats, "vendor": vend,
        })
    return pages


def summarize_pdf(pdf_name: str, pages: list, preds: list) -> dict:
    split_points = []
    for page, pred in zip(pages, preds):
        if pred["decision"] == "Break":
            split_points.append({"page": page["page_number"],
                                  "vendor": (page["vendor"] or {}).get("vendor_name")})
    distinct = []
    for sp in split_points:
        if sp["vendor"] and sp["vendor"] not in distinct:
            distinct.append(sp["vendor"])
    return {
        "pdf_name": pdf_name,
        "page_count": len(pages),
        "split_points": split_points,
        "invoice_count": len(split_points),
        "distinct_vendors": distinct,
        "is_mixed_vendor": len(distinct) > 1,
        "proposed_name": None,
    }


# ===========================================================================
# SECTION 8 — REPORT
# Plain English: write the Excel a human reviews. Sheet 1 "Summary" = one row
# per PDF with the predicted split pages and a DROPDOWN to judge it. Sheet 2
# "Page Detail" = one row per page explaining WHY (which rules fired). Styling
# is clean and lightly colored; the dropdown uses Excel data validation.
# ===========================================================================

def _format_splits(summary: dict) -> str:
    parts = []
    for sp in summary["split_points"]:
        parts.append(f"p{sp['page']} - {sp['vendor']}" if summary["is_mixed_vendor"] and sp["vendor"]
                     else f"p{sp['page']}")
    return ", ".join(parts)


def build_excel_report(summaries: list, page_details: list, out_path: Path, config: dict) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    # Open the EXISTING workbook if it's already there, so any other tabs the
    # reviewer made (e.g. a template sheet) are preserved. Only the two tabs we
    # regenerate below get refreshed. If there's no file yet, start a fresh one.
    if out_path.exists():
        wb = load_workbook(out_path)
    else:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        wb.remove(wb.active)  # drop the default empty "Sheet"

    # ---- Sheet: Summary  (rebuilt fresh each run, in the reviewer's 8-col format)
    if "Summary" in wb.sheetnames:
        del wb["Summary"]
    ws = wb.create_sheet("Summary", 0)   # keep Summary as the first tab
    cols = ["PDF Original Name", "Proposed PDF Name", "Pages", "New invoice starts at",
            "# invoices", "Vendors", "Reviewer verdict", "Correct split pages (if wrong)"]
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = header_fill, header_font

    verdicts = '"✅ Correct,⚠️ Partially correct,❌ Incorrect,🔍 Needs a closer look"'
    dv = DataValidation(type="list", formula1=verdicts, allow_blank=True)
    ws.add_data_validation(dv)

    for row_i, s in enumerate(summaries, start=2):
        ws.append([
            s["pdf_name"],                                          # A  original name (untouched)
            s.get("proposed_name") or "",                          # B  proposed name — TEXT ONLY, no copy
            s["page_count"],                                       # C  pages
            _format_splits(s),                                     # D  new invoice starts at
            s["invoice_count"],                                    # E  # invoices
            ", ".join(s["distinct_vendors"]) or "(none detected)", # F  vendors
            "",                                                    # G  reviewer verdict (dropdown)
            "",                                                    # H  correct split pages (if wrong)
        ])
        dv.add(ws.cell(row=row_i, column=7))  # dropdown on "Reviewer verdict" (col G)
    for width, col in zip([30, 30, 7, 30, 10, 24, 22, 26], "ABCDEFGH"):
        ws.column_dimensions[col].width = width

    # ---- Sheet: Page Detail  (rebuilt fresh each run, kept right after Summary)
    if "Page Detail" in wb.sheetnames:
        del wb["Page Detail"]
    wd = wb.create_sheet("Page Detail", 1)
    dcols = ["PDF", "Page", "Decision", "Confidence", "Score", "Why (rules fired)", "Read issue"]
    wd.append(dcols)
    for c in range(1, len(dcols) + 1):
        cell = wd.cell(row=1, column=c)
        cell.fill, cell.font = header_fill, header_font
    green = PatternFill("solid", fgColor="C6EFCE")
    for d in page_details:
        page, pred = d["page"], d["pred"]
        wd.append([
            d["pdf_name"], pred["page_number"], pred["decision"], pred["confidence"],
            pred["total_score"], "; ".join(pred["fired_rules"]),
            page.get("read_reason") or "",
        ])
        if pred["decision"] == "Break":
            wd.cell(row=wd.max_row, column=3).fill = green
    for cell in wd["F"]:
        cell.alignment = wrap
    for width, col in zip([30, 6, 10, 12, 8, 60, 30], "ABCDEFG"):
        wd.column_dimensions[col].width = width

    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(out_path)
    except PermissionError:
        raise SystemExit(
            f"\nCould not save {out_path.name} — it looks like the file is open in Excel.\n"
            f"Please close it and run the script again."
        )


# ===========================================================================
# SECTION 7 — PROPOSED NAME
# Plain English: work out what a single-vendor PDF WOULD be renamed to and put
# that name (as plain TEXT) in the Excel's "Proposed PDF Name" column. We do
# NOT create any renamed copy — nothing is written to disk here. If the same
# vendor appears again we add _01, _02, ... A PDF with MULTIPLE vendors (or no
# vendor detected) just gets a short note instead of a name.
# ===========================================================================

_vendor_counts: dict = {}

def reset_vendor_counts():
    _vendor_counts.clear()

def _dedup_name(base: str) -> str:
    n = _vendor_counts.get(base, 0)
    _vendor_counts[base] = n + 1
    return base if n == 0 else f"{base}_{n:02d}"


def assign_proposed_name(summary: dict) -> None:
    # Text only — no file is ever copied or renamed.
    if summary["is_mixed_vendor"]:
        summary["proposed_name"] = "(mixed vendors)"
        return
    if not summary["distinct_vendors"]:
        summary["proposed_name"] = "(no vendor detected)"
        return
    vendor = summary["distinct_vendors"][0]
    new_stem = _dedup_name(sanitize_filename(vendor))
    summary["proposed_name"] = f"{new_stem}.pdf"


# ===========================================================================
# SECTION 9 — SCORER  (run with --score, AFTER a human fills the Excel)
# Plain English: compare what the script predicted against what the reviewer
# said was correct, and report accuracy (Precision/Recall/F1). If the reviewer
# marked a row "Correct", the predicted split pages ARE the truth. Otherwise we
# use the pages they typed in "Correct split pages (if wrong)".
# ===========================================================================

def _parse_pages(cell: str) -> set:
    if not cell:
        return set()
    return {int(n) for n in re.findall(r"\d+", str(cell))}

def _prf(tp: int, fp: int, fn: int) -> dict:
    precision = tp / (tp + fp) if (tp + fp) else 0.0
    recall = tp / (tp + fn) if (tp + fn) else 0.0
    f1 = (2 * precision * recall / (precision + recall)) if (precision + recall) else 0.0
    return {"true_positives": tp, "false_positives": fp, "false_negatives": fn,
            "precision": precision, "recall": recall, "f1": f1}


def score_reviewed_report(xlsx_path: Path) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path)
    ws = wb["Summary"]
    tp = fp = fn = pdfs = pred_total = true_total = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, ""):
            continue
        verdict = (row[6] or "").strip()             # "Reviewer verdict" (col G)
        if not verdict:
            continue  # not yet reviewed — skip rather than count as wrong
        pdfs += 1
        predicted = _parse_pages(row[3])            # "New invoice starts at" (col D)
        correction = _parse_pages(row[7])            # "Correct split pages (if wrong)" (col H)
        truth = predicted if "Correct" in verdict and "Partially" not in verdict else correction
        pred_total += len(predicted)
        true_total += len(truth)
        tp += len(predicted & truth)
        fp += len(predicted - truth)
        fn += len(truth - predicted)

    metrics = _prf(tp, fp, fn)
    metrics.update({"pdfs": pdfs, "boundaries_predicted": pred_total,
                    "boundaries_true": true_total})
    return metrics


# ===========================================================================
# SELFTEST
# Plain English: built-in checks that prove the important logic still works.
# Run them any time with: python invoice_boundary_detector.py --selftest
# ===========================================================================

def _check_load_config():
    cfg = load_config("config.json")
    assert cfg["thresholds"]["break_at_score"] == 50, "break threshold should load from config"
    assert cfg["rule_weights"]["invoice_number_changed"] == 100, "weights should load from config"
    assert isinstance(cfg["known_vendors"], list), "known_vendors must be a list"

def _check_collect_pdfs_sorted(tmp=SCRIPT_DIR):
    # Uses whatever PDFs exist in input/; just verifies the function returns a
    # sorted list of Path objects ending in .pdf (empty is acceptable).
    cfg = load_config("config.json")
    pdfs = collect_pdfs(SCRIPT_DIR / cfg["paths"]["input_folder"])
    assert isinstance(pdfs, list), "collect_pdfs must return a list"
    assert all(str(p).lower().endswith(".pdf") for p in pdfs), "only .pdf paths"
    assert pdfs == sorted(pdfs), "collect_pdfs must be sorted"

def _check_invoice_number():
    assert find_doc_number("Invoice #: INV-1001") == "INV-1001"
    assert find_doc_number("no numbers here") is None

def _check_total_and_date():
    f = extract_features("Invoice No: 77\nDate: 03/14/2026\nTotal Due: $1,234.56")
    assert f["invoice_number"] == "77", f["invoice_number"]
    assert f["date"] == "03/14/2026", f["date"]
    assert f["total"] == "1,234.56", f["total"]

def _check_vendor_known_list():
    cfg = {"known_vendors": ["Ferguson Facility Supply"]}
    v = detect_vendor("Charges from Ferguson Facility Supply for July", cfg)
    assert v["vendor_name"] == "Ferguson Facility Supply"
    assert v["confidence"] == "high" and v["source"] == "known-list"

def _check_vendor_domain():
    cfg = {"known_vendors": []}
    v = detect_vendor("Questions? billing@acmeplumbing.com", cfg)
    assert v["vendor_name"] and "acmeplumbing" in v["vendor_name"].lower()
    assert v["source"] == "domain"

def _check_sanitize_filename():
    assert sanitize_filename('A/B:C*?"') == "A_B_C___"

def _mk_page(n, text="", inv=None, vendor=None):
    return {"page_number": n, "text": text,
            "features": {"invoice_number": inv},
            "vendor": {"vendor_name": vendor}}

def _check_invoice_number_changed_breaks():
    cfg = load_config("config.json")
    prev = _mk_page(1, inv="1001", vendor="Acme")
    curr = _mk_page(2, inv="1002", vendor="Acme")
    pred = score_boundary(prev, curr, cfg)
    assert pred["decision"] == "Break", pred
    assert any("Invoice number changed" in r for r in pred["fired_rules"]), pred

def _check_continuation_continues():
    cfg = load_config("config.json")
    prev = _mk_page(1, text="Invoice # 1001", inv="1001", vendor="Acme")
    curr = _mk_page(2, text="continued... page 2 of 3", inv=None, vendor="Acme")
    pred = score_boundary(prev, curr, cfg)
    assert pred["decision"] == "Continue", pred

def _check_first_page_always_break():
    cfg = load_config("config.json")
    preds = predict_boundaries([_mk_page(1, inv="1001")], cfg)
    assert preds[0]["decision"] == "Break"

def _check_summary_counts_splits():
    cfg = load_config("config.json")
    pages = [
        {"page_number": 1, "text": "Invoice # 1", "features": {"invoice_number": "1"}, "vendor": {"vendor_name": "Acme"}},
        {"page_number": 2, "text": "continued", "features": {"invoice_number": None}, "vendor": {"vendor_name": "Acme"}},
        {"page_number": 3, "text": "Invoice # 2", "features": {"invoice_number": "2"}, "vendor": {"vendor_name": "Beta"}},
    ]
    preds = predict_boundaries(pages, cfg)
    s = summarize_pdf("x.pdf", pages, preds)
    assert s["invoice_count"] == 2, s
    assert [sp["page"] for sp in s["split_points"]] == [1, 3], s
    assert s["is_mixed_vendor"] is True, s

def _check_dedup_name():
    reset_vendor_counts()
    assert _dedup_name("Acme") == "Acme"
    assert _dedup_name("Acme") == "Acme_01"
    assert _dedup_name("Acme") == "Acme_02"
    assert _dedup_name("Beta") == "Beta"

def _check_parse_pages():
    assert _parse_pages("1, 5, 9") == {1, 5, 9}
    assert _parse_pages("p1, p4") == {1, 4}
    assert _parse_pages("") == set()

def _check_prf_math():
    # tp=2, fp=1, fn=1 -> precision=2/3, recall=2/3, f1=2/3
    m = _prf(tp=2, fp=1, fn=1)
    assert round(m["precision"], 3) == 0.667 and round(m["recall"], 3) == 0.667
    assert round(m["f1"], 3) == 0.667

def _check_score_skips_unreviewed_rows():
    # A reviewer who only marks SOME rows must not have the blank rows
    # counted as "completely wrong" — they should be skipped entirely.
    import tempfile
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["PDF Original Name", "Proposed PDF Name", "Pages", "New invoice starts at",
                "# invoices", "Vendors", "Reviewer verdict", "Correct split pages (if wrong)"])
    ws.append(["a.pdf", "Acme.pdf", 3, "p1, p3", 2, "Acme", "✅ Correct", ""])
    ws.append(["b.pdf", "Acme_01.pdf", 3, "p1, p2, p3", 3, "Acme", "", ""])  # unreviewed
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "test.xlsx"
        wb.save(path)
        m = score_reviewed_report(path)
    assert m["pdfs"] == 1, m
    assert m["true_positives"] == 2 and m["false_positives"] == 0, m

SELFTESTS = [_check_load_config, _check_collect_pdfs_sorted, _check_invoice_number, _check_total_and_date,
             _check_vendor_known_list, _check_vendor_domain, _check_sanitize_filename,
             _check_invoice_number_changed_breaks, _check_continuation_continues, _check_first_page_always_break,
             _check_summary_counts_splits, _check_dedup_name, _check_parse_pages, _check_prf_math,
             _check_score_skips_unreviewed_rows]

def run_selftests() -> int:
    failed = 0
    for check in SELFTESTS:
        try:
            check()
            print(f"PASS  {check.__name__}")
        except AssertionError as e:
            failed += 1
            print(f"FAIL  {check.__name__}: {e}")
    print(f"\n{len(SELFTESTS) - failed}/{len(SELFTESTS)} checks passed.")
    return 1 if failed else 0


# ===========================================================================
# SECTION 10 — MAIN
# Plain English: the front door. It looks at the command-line flag and decides
# which mode to run.
# ===========================================================================

def main():
    parser = argparse.ArgumentParser(description="Invoice boundary detection (no splitting).")
    parser.add_argument("--score", action="store_true", help="Score a reviewed Excel file.")
    parser.add_argument("--selftest", action="store_true", help="Run built-in logic checks.")
    args = parser.parse_args()

    if args.selftest:
        sys.exit(run_selftests())

    if args.score:
        config = load_config()
        xlsx = SCRIPT_DIR / config["paths"]["output_reports"] / "predictions.xlsx"
        if not xlsx.exists():
            print(f"No report at {xlsx}. Run prediction mode first, then fill the Excel.")
            return
        m = score_reviewed_report(xlsx)
        print("Accuracy from reviewed report:")
        print(f"  PDFs reviewed        : {m['pdfs']}")
        print(f"  Boundaries predicted : {m['boundaries_predicted']}")
        print(f"  Boundaries (truth)   : {m['boundaries_true']}")
        print(f"  True positives       : {m['true_positives']}")
        print(f"  False positives      : {m['false_positives']}")
        print(f"  False negatives      : {m['false_negatives']}")
        print(f"  Precision            : {m['precision']:.3f}")
        print(f"  Recall               : {m['recall']:.3f}")
        print(f"  F1 score             : {m['f1']:.3f}")
        return

    # Default = prediction mode (built up over the walking-skeleton tasks).
    config = load_config()
    input_folder = SCRIPT_DIR / config["paths"]["input_folder"]
    pdfs = collect_pdfs(input_folder)
    if not pdfs:
        print(f"No PDFs found in {input_folder}. Put sample PDFs there and re-run.")
        return

    reset_vendor_counts()
    all_summaries, all_page_details = [], []
    for pdf_path in tqdm(pdfs, desc="PDFs", unit="file"):
        records = read_pdf_pages(pdf_path, config)
        pages = build_pages(records, config)
        preds = predict_boundaries(pages, config)
        summary = summarize_pdf(pdf_path.name, pages, preds)
        assign_proposed_name(summary)
        all_summaries.append(summary)
        for page, pred in zip(pages, preds):
            all_page_details.append({"pdf_name": pdf_path.name, "page": page, "pred": pred})
        splits = ", ".join(str(sp["page"]) for sp in summary["split_points"])
        print(f"{pdf_path.name}: {summary['invoice_count']} invoice(s), splits at [{splits}]")

    out_dir = SCRIPT_DIR / config["paths"]["output_reports"]
    out_dir.mkdir(parents=True, exist_ok=True)
    build_excel_report(all_summaries, all_page_details, out_dir / "predictions.xlsx", config)
    print(f"\nReport written to {out_dir / 'predictions.xlsx'}")


if __name__ == "__main__":
    main()
