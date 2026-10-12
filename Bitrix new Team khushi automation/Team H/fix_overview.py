import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
import pandas as pd
import re

file_name = 'Team_H_Detailed_Overdue_Tasks.xlsx'
orig_file = 'Team_H_Overdue_Tasks.xlsx'

print("Loading mapping...")
orig_overview = pd.read_excel(orig_file, sheet_name='Overview', skiprows=3)
l1_mapping = {}
for _, row in orig_overview.iterrows():
    code = row.get('Client Code')
    l1 = row.get('L1 (Project Lead)')
    if pd.notna(code) and pd.notna(l1):
        l1_mapping[str(code).strip()] = str(l1).strip()

l1_mapping['C149A'] = 'Zeel Patel'
l1_mapping['C283A'] = 'Krishnakumar Bang'
l1_mapping['C352A'] = 'Zeel Patel'

print(f"Modifying {file_name}...")
wb = openpyxl.load_workbook(file_name)
ws = wb['Overview']

# The sheet currently has:
# Row 1-3: Titles
# Row 4: Project | Overdue Tasks Count
# Row 5+: Data

# We need to shift data and split 'Project' into 'Client Code', 'Client', 'L1 (Project Lead)'
pattern = r'^([A-Za-z0-9]+)\s*-\s*(.*?)\s*-\s*([^-]+?)\s*(?:-|\()\s*([A-Za-z0-9]+)\)?$'

new_data = []
# Start reading from row 4
if ws.cell(row=4, column=1).value == 'Client Code':
    print("Already formatted. Skipping.")
    exit(0)

# Start reading from row 5
for row in range(5, ws.max_row + 1):
    proj_val = ws.cell(row=row, column=1).value
    count_val = ws.cell(row=row, column=2).value
    if proj_val is None:
        continue
    
    match = re.match(pattern, proj_val)
    if match:
        _, client, _, code = match.groups()
        client_code = code.strip()
        client_name = client.strip()
        l1 = l1_mapping.get(client_code, '—')
    else:
        client_code = ''
        client_name = proj_val
        l1 = '—'
        
    new_data.append((client_code, client_name, l1, count_val))

# Clear existing data in rows 4+ 
ws.delete_rows(4, ws.max_row)

# Add new headers
ws.cell(row=4, column=1, value='Client Code')
ws.cell(row=4, column=2, value='Client')
ws.cell(row=4, column=3, value='L1 (Project Lead)')
ws.cell(row=4, column=4, value='Overdue Tasks')

# Add new data
for idx, data_row in enumerate(new_data):
    row_num = 5 + idx
    ws.cell(row=row_num, column=1, value=data_row[0])
    ws.cell(row=row_num, column=2, value=data_row[1])
    ws.cell(row=row_num, column=3, value=data_row[2])
    ws.cell(row=row_num, column=4, value=data_row[3])

# Apply styles
header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Light red/orange
red_font = Font(color="C00000", bold=True) # Dark red

for col in range(1, 5):
    c = ws.cell(row=4, column=col)
    c.fill = header_fill
    c.font = header_font

for row in range(5, 5 + len(new_data)):
    c = ws.cell(row=row, column=4)
    c.fill = red_fill
    c.font = red_font
    c.alignment = Alignment(horizontal='center')

# Add AutoFilter
ws.auto_filter.ref = f"A4:D{4 + len(new_data)}"

# Merge title row across the 4 columns
ws.merge_cells('A1:D1')

# Set column widths
ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 40
ws.column_dimensions['C'].width = 25
ws.column_dimensions['D'].width = 18

wb.save(file_name)
print("Overview sheet beautifully updated!")
