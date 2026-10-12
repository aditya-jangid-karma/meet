import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import glob

files = glob.glob('*_Live_Tasks.xlsx')

header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
title_font = Font(size=14, bold=True)
total_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
total_font = Font(bold=True)

red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
red_font = Font(color="C00000", bold=True)
yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
yellow_font = Font(color="B8860B", bold=True)
green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
green_font = Font(color="385723", bold=True)

thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'), 
    top=Side(style='thin'), bottom=Side(style='thin')
)
thick_border_top = Border(top=Side(style='medium', color="00B050")) # Greenish border in image
thick_border_bottom = Border(bottom=Side(style='medium', color="00B050"))
thick_border_left = Border(left=Side(style='medium', color="00B050"))
thick_border_right = Border(right=Side(style='medium', color="00B050"))

for f in files:
    print(f"Processing {f}...")
    wb = openpyxl.load_workbook(f)
    
    # --- OVERVIEW TAB ---
    ws_over = wb['Overview']
    max_r = ws_over.max_row
    
    sum_overdue = 0
    sum_today = 0
    sum_tomorrow = 0
    
    for r in range(5, max_r + 1):
        v1 = ws_over.cell(row=r, column=3).value or 0
        v2 = ws_over.cell(row=r, column=4).value or 0
        v3 = ws_over.cell(row=r, column=5).value or 0
        sum_overdue += int(v1)
        sum_today += int(v2)
        sum_tomorrow += int(v3)
        
    total_row = max_r + 1
    ws_over.cell(row=total_row, column=2, value="Total").font = total_font
    
    c_over = ws_over.cell(row=total_row, column=3, value=sum_overdue)
    c_over.fill = red_fill
    c_over.font = red_font
    c_over.alignment = Alignment(horizontal='center')
    
    c_tod = ws_over.cell(row=total_row, column=4, value=sum_today)
    c_tod.fill = yellow_fill
    c_tod.font = yellow_font
    c_tod.alignment = Alignment(horizontal='center')
    
    c_tom = ws_over.cell(row=total_row, column=5, value=sum_tomorrow)
    c_tom.fill = green_fill
    c_tom.font = green_font
    c_tom.alignment = Alignment(horizontal='center')
    
    ws_over.auto_filter.ref = f"A4:E{total_row}"

    # --- ALL CLIENT GROUPS TAB ---
    ws_all = wb.create_sheet('All Client Groups')
    
    current_row = 1
    
    sheet_names = wb.sheetnames
    
    for s_name in sheet_names:
        if s_name in ['Overview', 'All Client Groups']:
            continue
            
        # Parse project tasks
        ws_proj = wb[s_name]
        
        # Determine client name from sheet name (e.g. C134A FleetPeek -> FleetPeek)
        parts = s_name.split(' ', 1)
        client_nam = parts[1] if len(parts) > 1 else s_name
        
        # 1. Project Title Row
        ws_all.cell(row=current_row, column=1, value=s_name).font = title_font
        ws_all.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        current_row += 1
        
        # 2. Header Row
        headers = ['Task ID', 'Client Nam', 'Task Title', 'Deadline', 'Responsible Person', 'Status']
        for col_idx, h in enumerate(headers, 1):
            c = ws_all.cell(row=current_row, column=col_idx, value=h)
            c.fill = header_fill
            c.font = header_font
            
        current_row += 1
        
        # 3. Tasks
        for r in range(4, ws_proj.max_row + 1): # Tasks start at row 4 in old sheets
            task_id = ws_proj.cell(row=r, column=1).value
            if not task_id: continue
            
            ws_all.cell(row=current_row, column=1, value=task_id)
            ws_all.cell(row=current_row, column=2, value=client_nam)
            ws_all.cell(row=current_row, column=3, value=ws_proj.cell(row=r, column=2).value) # Title
            ws_all.cell(row=current_row, column=4, value=ws_proj.cell(row=r, column=3).value) # Deadline
            ws_all.cell(row=current_row, column=5, value=ws_proj.cell(row=r, column=4).value) # Resp
            ws_all.cell(row=current_row, column=6, value=ws_proj.cell(row=r, column=5).value) # Status
            
            # Alternate background shading just for better readability
            if (current_row % 2) == 0:
                for c_idx in range(1, 7):
                    ws_all.cell(row=current_row, column=c_idx).fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
                    
            current_row += 1
            
        current_row += 1 # Empty row between projects
        
        # Delete old project tab
        del wb[s_name]

    # Set column widths for All Client Groups
    ws_all.column_dimensions['A'].width = 12
    ws_all.column_dimensions['B'].width = 25
    ws_all.column_dimensions['C'].width = 60
    ws_all.column_dimensions['D'].width = 15
    ws_all.column_dimensions['E'].width = 25
    ws_all.column_dimensions['F'].width = 15
    
    out_name = f.replace('_Live_Tasks.xlsx', '_Final_Report.xlsx')
    wb.save(out_name)
    print(f"Saved {out_name}")

print("Done restructuring!")
