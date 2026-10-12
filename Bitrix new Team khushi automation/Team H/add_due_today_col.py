import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import re
from datetime import datetime
import urllib3
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

GROUPS_URL = "https://karma.bitrix24.in/rest/1/4ge56nup4fuub38j/sonet_group.get.json"
TASKS_URL = "https://karma.bitrix24.in/rest/1/euiyos3e551clsi5/tasks.task.list.json"
file_name = 'Team_H_Detailed_Overdue_Tasks.xlsx'

def get_due_today_counts():
    print("Fetching all projects to build mapping...")
    groups = []
    start = 0
    while True:
        resp = requests.get(f"{GROUPS_URL}?start={start}", verify=False).json()
        batch = resp.get('result', [])
        groups.extend(batch)
        if 'next' in resp:
            start = resp['next']
        else:
            break

    pattern = r'^([A-Za-z0-9]+)\s*-\s*(.*?)\s*-\s*([^-]+?)\s*(?:-|\()\s*([A-Za-z0-9]+)\)?$'
    
    # Map groupId (str) -> Client Code
    group_to_code = {}
    for g in groups:
        project_name = g.get('NAME', '')
        gid = str(g['ID'])
        match = re.match(pattern, project_name)
        if match:
            team, _, _, code = match.groups()
            if team.strip().upper() == 'H':
                group_to_code[gid] = code.strip()
                
    print("Fetching Tasks Due Today...")
    today = datetime.now().strftime('%Y-%m-%d')
    start_task = 0
    code_counts = {}
    
    while True:
        params = {
            'filter[>=DEADLINE]': f'{today}T00:00:00',
            'filter[<=DEADLINE]': f'{today}T23:59:59',
            'filter[!STATUS]': '5',
            'start': start_task
        }
        resp = requests.get(TASKS_URL, params=params, verify=False).json()
        
        result_node = resp.get('result', {})
        if isinstance(result_node, dict) and 'tasks' in result_node:
            batch = result_node['tasks']
        elif isinstance(result_node, list):
            batch = result_node
        else:
            batch = []
            
        for t in batch:
            g_id = str(t.get('groupId'))
            if g_id in group_to_code:
                c_code = group_to_code[g_id]
                code_counts[c_code] = code_counts.get(c_code, 0) + 1
                
        if 'next' in result_node:
            start_task = result_node['next']
        elif 'next' in resp:
            start_task = resp['next']
        else:
            if len(batch) == 50:
                start_task += 50
            else:
                break
                
    return code_counts

print("Modifying Excel file...")
code_counts = get_due_today_counts()

wb = openpyxl.load_workbook(file_name)
ws = wb['Overview']

# Add new Header
ws.cell(row=4, column=5, value='Due Today')

# Style the new header
header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
ws.cell(row=4, column=5).fill = header_fill
ws.cell(row=4, column=5).font = header_font

# Add the counts and styling for rows
yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Light yellow for Due Today
due_font = Font(color="B8860B", bold=True) # Dark Goldenrod

for row in range(5, ws.max_row + 1):
    c_code = str(ws.cell(row=row, column=1).value).strip()
    count = code_counts.get(c_code, 0)
    
    c = ws.cell(row=row, column=5, value=count)
    c.fill = yellow_fill
    c.font = due_font
    c.alignment = Alignment(horizontal='center')

# Merge title row across 5 columns now
ws.unmerge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)

# Adjust autofilter
ws.auto_filter.ref = f"A4:E{ws.max_row}"

# Set width for new column
ws.column_dimensions['E'].width = 15

wb.save(file_name)
print(f"Added 'Due Today' column to {file_name}!")
