import pandas as pd
import re
from datetime import datetime, timedelta
import requests
import urllib3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

TASKS_URL = "https://karma.bitrix24.in/rest/1/euiyos3e551clsi5/tasks.task.list.json"
GROUPS_URL = "https://karma.bitrix24.in/rest/1/4ge56nup4fuub38j/sonet_group.get.json"

# 1. Fetch Group Mapping to Client Code
print("Fetching groups to map IDs to Client Codes...")
groups = []
start = 0
while True:
    resp = requests.get(f"{GROUPS_URL}?start={start}", verify=False).json()
    batch = resp.get('result', [])
    groups.extend(batch)
    if 'next' in resp:
        start = resp['next']
    else:
        break

pattern = r'^([A-Za-z0-9]+)\s*-\s*(.*?)\s*-\s*([^-]+?)\s*(?:-|\()\s*([A-Za-z0-9]+)\)?$'
group_to_code = {}
for g in groups:
    project_name = g.get('NAME', '')
    gid = str(g['ID'])
    match = re.match(pattern, project_name)
    if match:
        team, _, _, code = match.groups()
        if team.strip().upper() == 'H':
            group_to_code[gid] = code.strip()

# 2. Fetch Tasks Due Tomorrow
print("Fetching Tasks Due Tomorrow...")
tomorrow = (datetime.now() + timedelta(days=1)).strftime('%Y-%m-%d')
start_task = 0
due_tomorrow_counts = {}

while True:
    params = {
        'filter[>=DEADLINE]': f'{tomorrow}T00:00:00',
        'filter[<=DEADLINE]': f'{tomorrow}T23:59:59',
        'filter[!STATUS]': '5',
        'start': start_task
    }
    resp = requests.get(TASKS_URL, params=params, verify=False).json()
    
    result_node = resp.get('result', {})
    if isinstance(result_node, dict) and 'tasks' in result_node:
        batch = result_node['tasks']
    elif isinstance(result_node, list):
        batch = result_node
    else:
        batch = []
        
    for t in batch:
        g_id = str(t.get('groupId'))
        # Only count if the creator/responsible isn't dismissed? Bitrix doesn't easily filter by "not contains dismissed".
        # We will just count them for now. If needed, we'd have to inspect the responsible person of the due tomorrow tasks.
        # But wait! We can just fetch them and check the responsible person name!
        resp_person = t.get('responsible', {})
        resp_name = resp_person.get('name', '').lower()
        if 'dismissed' not in resp_name:
            if g_id in group_to_code:
                c_code = group_to_code[g_id]
                due_tomorrow_counts[c_code] = due_tomorrow_counts.get(c_code, 0) + 1
            
    if 'next' in result_node:
        start_task = result_node['next']
    elif 'next' in resp:
        start_task = resp['next']
    else:
        if len(batch) == 50:
            start_task += 50
        else:
            break

# 3. Load Excel Files
master_detailed = 'Team_H_Detailed_Overdue_Tasks.xlsx'
orig_mapping = 'Team_H_Overdue_Tasks.xlsx'

print("Loading Excel mapping...")
orig_overview = pd.read_excel(orig_mapping, sheet_name='Overview', skiprows=3)
l1_mapping = {}
client_mapping = {}
for _, row in orig_overview.iterrows():
    code = row.get('Client Code')
    l1 = row.get('L1 (Project Lead)')
    client = row.get('Client')
    if pd.notna(code):
        c = str(code).strip()
        l1_mapping[c] = str(l1).strip() if pd.notna(l1) else 'Unknown'
        client_mapping[c] = str(client).strip()

l1_mapping['C149A'] = 'Zeel Patel'
l1_mapping['C283A'] = 'Krishnakumar Bang'
l1_mapping['C352A'] = 'Zeel Patel'

print("Processing master sheets...")
xl = pd.ExcelFile(master_detailed)
overview_df = xl.parse('Overview', skiprows=3)

# 4. Split by L1 and generate files
l1_groups = {}
# Find all unique L1s
for code, l1 in l1_mapping.items():
    if l1 != 'Unknown' and l1 != '—' and str(l1).strip() != '':
        if l1 not in l1_groups:
            l1_groups[l1] = []
        l1_groups[l1].append(code)

for l1, codes in l1_groups.items():
    safe_l1 = re.sub(r'[^a-zA-Z0-9_\- ]', '', l1).strip().replace(' ', '_')
    out_file = f"{safe_l1}_Live_Tasks.xlsx"
    print(f"Generating {out_file}...")
    
    # We will rebuild the overview data for this L1
    l1_overview_data = []
    project_sheets = {}
    
    for code in codes:
        client_name = client_mapping.get(code, code)
        
        # Look for a sheet matching this project
        matching_sheet = None
        for sheet_name in xl.sheet_names:
            if sheet_name == 'Overview': continue
            if code in sheet_name:
                matching_sheet = sheet_name
                break
        
        overdue_count = 0
        df = pd.DataFrame()
        
        if matching_sheet:
            # We don't skip rows because team_h_overdue_tasks.py writes headers on row 1!
            df = xl.parse(matching_sheet)
            if 'Responsible Person' in df.columns:
                df = df[~df['Responsible Person'].str.lower().str.contains('dismissed', na=False)]
            overdue_count = len(df)
            
        due_today = 0
        # We can extract Due Today from the master overview if it exists
        # Find the row in master overview
        master_row = overview_df[overview_df['Client Code'] == code]
        if not master_row.empty and 'Due Today' in master_row.columns:
            due_today = master_row.iloc[0]['Due Today']
            if pd.isna(due_today): due_today = 0
            
        due_tomorrow = due_tomorrow_counts.get(code, 0)
        
        # Only include in Overview if it has at least some tasks
        if overdue_count > 0 or due_today > 0 or due_tomorrow > 0:
            l1_overview_data.append({
                'Client Code': code,
                'Client': client_name,
                'df': df if matching_sheet else pd.DataFrame(),
                'Overdue Tasks': overdue_count,
                'Due Today': int(due_today),
                'Due Tomorrow': due_tomorrow
            })
            
            if overdue_count > 0:
                project_sheets[matching_sheet] = df

    if not l1_overview_data:
        print(f"No active tasks for {l1}. Skipping.")
        continue
        
    # Sort by Overdue Tasks first
    l1_overview_data.sort(key=lambda x: x['Overdue Tasks'], reverse=True)
    
    # We will build the layout the user requested: one row per client.
    # We will not include the L1 column, since the whole file is for this L1.
    expanded_data = []
    for item in l1_overview_data:
        expanded_data.append({
            'Client Code': item['Client Code'],
            'Client': item['Client'],
            'Overdue Tasks': item['Overdue Tasks'],
            'Due Today': item['Due Today'],
            'Due Tomorrow': item['Due Tomorrow']
        })
        
    overview_df_filtered = pd.DataFrame(expanded_data)
    
    with pd.ExcelWriter(out_file, engine='openpyxl') as writer:
        overview_df_filtered.to_excel(writer, sheet_name='Overview', index=False, startrow=3)
        
        for sheet_name, df_proj in project_sheets.items():
            df_proj.to_excel(writer, sheet_name=sheet_name, index=False, startrow=2)
            
    # Apply styling
    wb = openpyxl.load_workbook(out_file)
    ws = wb['Overview']
    
    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    red_font = Font(color="C00000", bold=True)
    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    yellow_font = Font(color="B8860B", bold=True)
    green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    green_font = Font(color="385723", bold=True)
    
    # Add title rows
    ws['A1'] = f"{l1} - Overdue & Upcoming Tasks"
    ws['A1'].font = Font(size=14, bold=True)
    ws.merge_cells('A1:E1')
    ws['A2'] = "Filtered to exclude dismissed users"
    ws['A2'].font = Font(size=11, italic=True)
    ws['A3'] = "Includes Due Today and Due Tomorrow counts"
    ws['A3'].font = Font(size=11, italic=True)
    
    for col in range(1, 6):
        c = ws.cell(row=4, column=col)
        c.fill = header_fill
        c.font = header_font
        
    for row in range(5, ws.max_row + 1):
        client_code_val = ws.cell(row=row, column=1).value
        if not client_code_val:
            continue
            
        # Overdue
        c_over = ws.cell(row=row, column=3)
        c_over.fill = red_fill
        c_over.font = red_font
        c_over.alignment = Alignment(horizontal='center')
        # Due Today
        c_today = ws.cell(row=row, column=4)
        c_today.fill = yellow_fill
        c_today.font = yellow_font
        c_today.alignment = Alignment(horizontal='center')
        # Due Tomorrow
        c_tom = ws.cell(row=row, column=5)
        c_tom.fill = green_fill
        c_tom.font = green_font
        c_tom.alignment = Alignment(horizontal='center')
        
    ws.auto_filter.ref = f"A4:E{ws.max_row}"
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 18
    ws.freeze_panes = 'A5'
    
    for sheet_name in wb.sheetnames:
        if sheet_name == 'Overview': continue
        ws_proj = wb[sheet_name]
        ws_proj['A1'] = sheet_name
        ws_proj['A1'].font = Font(size=14, bold=True)
        ws_proj.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ws_proj.max_column)
        
        for col in range(1, ws_proj.max_column + 1):
            c = ws_proj.cell(row=3, column=col)
            c.fill = header_fill
            c.font = header_font
            
        ws_proj.column_dimensions['A'].width = 12
        ws_proj.column_dimensions['B'].width = 60
        ws_proj.column_dimensions['C'].width = 15
        ws_proj.column_dimensions['D'].width = 25
        ws_proj.column_dimensions['E'].width = 20
        ws_proj.column_dimensions['F'].width = 15
        ws_proj.freeze_panes = 'A4'
        
    wb.save(out_file)

print("Done generating fresh split files!")
