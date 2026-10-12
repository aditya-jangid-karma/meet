import os
import glob
import re
import requests
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
import urllib3
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

BATCH_URL = "https://karma.bitrix24.in/rest/1/euiyos3e551clsi5/batch.json"
GROUPS_URL = "https://karma.bitrix24.in/rest/1/4ge56nup4fuub38j/sonet_group.get.json"

# Fetch all groups to get IDs by name quickly
print("Fetching all projects to build mapping for Due Today...")
all_groups = []
start = 0
while True:
    resp = requests.get(f"{GROUPS_URL}?start={start}", verify=False).json()
    batch = resp.get('result', [])
    all_groups.extend(batch)
    if 'next' in resp:
        start = resp['next']
    else:
        break

project_name_to_id = {g.get('NAME', '').strip(): str(g['ID']) for g in all_groups}

teams_dir = r"C:\Users\Meet Patel\Desktop\Bitrix new Team khushi automation\teams"
files = glob.glob(os.path.join(teams_dir, 'Team_*_Detailed_Overdue_Tasks.xlsx'))

pattern = r'^([A-Za-z0-9]+)\s*-\s*(.*?)\s*-\s*([^-]+?)\s*(?:-|\()\s*([A-Za-z0-9]+)\)?$'

today_str = datetime.now().strftime('%Y-%m-%d')
start_of_today = f"{today_str}T00:00:00"
end_of_today = f"{today_str}T23:59:59"

for file_path in files:
    print(f"\nProcessing {os.path.basename(file_path)}...")
    wb = openpyxl.load_workbook(file_path)
    ws = wb['Overview']
    
    # Check if already formatted
    if ws.cell(row=4, column=1).value == 'Client Code':
        print("Already formatted. Skipping.")
        continue
        
    # First, read existing data
    new_data = []
    projects_to_fetch = []
    
    for row in range(1, ws.max_row + 1):
        proj_val = ws.cell(row=row, column=1).value
        count_val = ws.cell(row=row, column=2).value
        
        # Skip if it's a message row or header row
        if not proj_val or proj_val == 'Project' or str(proj_val).startswith('No overdue'):
            continue
            
        proj_str = str(proj_val).strip()
        gid = project_name_to_id.get(proj_str)
        if gid:
            projects_to_fetch.append((proj_str, gid, count_val))
        else:
            print(f"Warning: Could not find ID for project: {proj_str}")
            projects_to_fetch.append((proj_str, None, count_val))

    # Fetch Due Today counts
    due_today_counts = {}
    if projects_to_fetch:
        gids_to_fetch = [p[1] for p in projects_to_fetch if p[1] is not None]
        for i in range(0, len(gids_to_fetch), 50):
            chunk = gids_to_fetch[i:i+50]
            cmd = {}
            for gid in chunk:
                cmd[f'g_{gid}'] = f'tasks.task.list?filter[GROUP_ID]={gid}&filter[>=DEADLINE]={start_of_today}&filter[<=DEADLINE]={end_of_today}&filter[!STATUS]=5'
            
            batch_resp = requests.post(BATCH_URL, json={'halt': 0, 'cmd': cmd}, verify=False).json()
            results = batch_resp.get('result', {}).get('result', {})
            
            for k, v in results.items():
                gid = k.replace('g_', '')
                tasks = v.get('tasks', [])
                due_today_counts[gid] = len(tasks)

    # Build rows
    for proj_str, gid, count_val in projects_to_fetch:
        match = re.match(pattern, proj_str)
        if match:
            _, client, person, code = match.groups()
            client_code = code.strip()
            client_name = client.strip()
            l1 = person.strip() # We don't have mapping, so use person from string
        else:
            client_code = ''
            client_name = proj_str
            l1 = ''
            
        due_today = due_today_counts.get(gid, 0) if gid else 0
        new_data.append((client_code, client_name, l1, count_val, due_today))
        
    # Clear the entire sheet safely
    ws.delete_rows(1, ws.max_row)
        
    # Insert blank rows at top to match layout
    ws.append([])
    ws.append([])
    ws.append([])
    
    # Add Headers
    ws.append(['Client Code', 'Client', 'L1 (Project Lead)', 'Overdue Tasks', 'Due Today'])
    
    # Add Data
    for data_row in new_data:
        ws.append(data_row)
        
    # Styling
    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    red_font = Font(color="C00000", bold=True)
    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    yellow_font = Font(color="B8860B", bold=True)
    
    for col in range(1, 6):
        c = ws.cell(row=4, column=col)
        c.fill = header_fill
        c.font = header_font
        
    for row in range(5, 5 + len(new_data)):
        c_over = ws.cell(row=row, column=4)
        c_over.fill = red_fill
        c_over.font = red_font
        c_over.alignment = Alignment(horizontal='center')
        
        c_today = ws.cell(row=row, column=5)
        c_today.fill = yellow_fill
        c_today.font = yellow_font
        c_today.alignment = Alignment(horizontal='center')

    if new_data:
        ws.auto_filter.ref = f"A4:E{4 + len(new_data)}"
        
    ws.merge_cells('A1:E1')
    
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 18

    wb.save(file_path)
    print("Done!")

print("\nAll team files formatted perfectly!")
